"""Report sheets: PNL, BALANCE_SHEET, CASH_FLOW, MONTHLY_SUMMARY, TRIAL_BALANCE, LEDGER.

Reports never hold typed numbers. Detail lines are "slots" filled from CALC by
(type, rank); section totals are computed straight from CALC by type, so a
report stays correct even if a user adds more accounts than there are slots
(a check line says so).
"""

from __future__ import annotations

from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

from . import model as M
from . import styles as S
from .xl import add_date_validation, add_list_validation, define, protect, status_colors, title_block, unlock

PERIOD_TEXT = '="For the month of "&TEXT(RptStart,"mmmm yyyy")&"  ·  Fiscal year "&TEXT(FYStart,"mmm yyyy")&" – "&TEXT(FYEnd,"mmm yyyy")&"  ·  Amounts in "&Currency'


def _report_header(ws, title: str, period_formula: str, width: int):
    title_block(ws, title, "", width)
    ws["B2"].value = '=BizName'
    ws["B2"].font = S.font(size=11, bold=True, color=S.SLATE)
    ws["B4"].value = period_formula
    ws["B4"].font = S.font(size=9, italic=True, color=S.MUTED)


def _colhead(ws, row, col, text, sub=None, fmt=None):
    c = ws.cell(row=row, column=col, value=text)
    c.font = S.font(bold=True, color=S.WHITE, size=9)
    c.fill = S.fill(S.NAVY)
    c.alignment = S.CENTER
    if fmt:
        c.number_format = fmt
    return c


def _section(ws, row, label, ncols):
    c = ws.cell(row=row, column=2, value=label)
    c.font = S.font(bold=True, color=S.TEAL)
    for col in range(2, 2 + ncols):
        ws.cell(row=row, column=col).border = S.Border(bottom=S.Side(style="thin", color=S.TEAL))


def _total(ws, row, label, ncols, strong=False, fmt=S.NUM):
    c = ws.cell(row=row, column=2, value=label)
    c.font = S.font(bold=True, color=S.NAVY if strong else S.SLATE, size=11 if strong else 10)
    for col in range(3, 2 + ncols):
        cell = ws.cell(row=row, column=col)
        cell.font = S.font(bold=True, color=S.NAVY if strong else S.SLATE)
        cell.number_format = fmt
        cell.border = S.TOTAL_RULE if strong else S.TOP_RULE
        if strong:
            cell.fill = S.fill(S.LIGHT)
    if strong:
        c.fill = S.fill(S.LIGHT)


def _print_setup(ws, landscape=True, area=None):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    if area:
        ws.print_area = area


# ----------------------------------------------------------------------------- P&L
PL_COLS = [("C", "calc_Cur"), ("D", "calc_Prev"), ("E", "calc_YTD"), ("F", "calc_FY")] + [
    (get_column_letter(8 + k), f"calc_M{k + 1}") for k in range(12)]


def build_pnl(wb, ws) -> dict:
    _report_header(ws, "Profit & Loss", PERIOD_TEXT, 18)
    ws.column_dimensions["B"].width = 34
    for col, _ in PL_COLS:
        ws.column_dimensions[col].width = 13
    ws.column_dimensions["G"].width = 2
    hr = 6
    _colhead(ws, hr, 2, "Account")
    for col, label in [("C", "This Month"), ("D", "Previous Month"), ("E", "Year to Date"), ("F", "Full Year")]:
        _colhead(ws, hr, ws[f"{col}1"].column, label)
    for k in range(12):
        c = _colhead(ws, hr, 8 + k, f"=EDATE(FYStart,{k})", fmt="mmm yy")
    ws.row_dimensions[hr].height = 28
    rows = {}
    r = hr + 2

    def line_values(row, sign):
        for col, rng in PL_COLS:
            cell = ws[f"{col}{row}"]
            cell.value = f'=IF($B{row}="","",{sign}*IFERROR(INDEX({rng},MATCH($B{row},calc_Name,0)),0))'
            cell.number_format = S.NUM
            cell.font = S.font()

    def total_values(row, expr_by_rng):
        for col, rng in PL_COLS:
            ws[f"{col}{row}"].value = "=" + expr_by_rng(rng, col)

    def section(typ, label, sign):
        nonlocal r
        _section(ws, r, label.upper(), 18)
        r += 1
        for k in range(1, M.PL_SLOTS[typ] + 1):
            c = ws.cell(row=r, column=2, value=f'=IFERROR(INDEX(calc_Name,MATCH("{typ}|{k}",calc_Key,0)),"")')
            c.alignment = S.INDENT1
            c.font = S.font()
            line_values(r, sign)
            r += 1
        _total(ws, r, f"Total {label}", 18)
        total_values(r, lambda rng, col: f'{sign}*SUMIFS({rng},calc_Type,"{typ}")')
        rows[typ] = r
        r += 2

    section("Revenue", "Revenue", -1)
    section("Cost of Sales", "Cost of Sales", 1)
    _total(ws, r, "GROSS PROFIT", 18, strong=True)
    total_values(r, lambda rng, col: f'{col}{rows["Revenue"]}-{col}{rows["Cost of Sales"]}')
    rows["GP"] = r
    r += 1
    ws.cell(row=r, column=2, value="Gross margin").font = S.font(italic=True, color=S.MUTED)
    for col, _ in PL_COLS:
        c = ws[f"{col}{r}"]
        c.value = f'=IF({col}{rows["Revenue"]}=0,"",{col}{rows["GP"]}/{col}{rows["Revenue"]})'
        c.number_format = S.PCT
        c.font = S.font(italic=True, color=S.MUTED)
    rows["GM"] = r
    r += 2
    section("Operating Expense", "Operating Expenses", 1)
    _total(ws, r, "NET OPERATING INCOME", 18, strong=True)
    total_values(r, lambda rng, col: f'{col}{rows["GP"]}-{col}{rows["Operating Expense"]}')
    rows["NOI"] = r
    r += 2
    section("Other Income", "Other Income", -1)
    section("Other Expense", "Other Expenses", 1)
    _total(ws, r, "NET PROFIT", 18, strong=True)
    total_values(r, lambda rng, col: f'{col}{rows["NOI"]}+{col}{rows["Other Income"]}-{col}{rows["Other Expense"]}')
    rows["NET"] = r
    r += 1
    ws.cell(row=r, column=2, value="Net margin").font = S.font(italic=True, color=S.MUTED)
    for col, _ in PL_COLS:
        c = ws[f"{col}{r}"]
        c.value = f'=IF({col}{rows["Revenue"]}=0,"",{col}{rows["NET"]}/{col}{rows["Revenue"]})'
        c.number_format = S.PCT
        c.font = S.font(italic=True, color=S.MUTED)
    rows["NM"] = r
    r += 2
    hidden = "+".join(f'MAX(0,COUNTIFS(calc_Type,"{t}",calc_Name,"<>")-{n})' for t, n in M.PL_SLOTS.items())
    ws.cell(row=r, column=2, value=f'=IF(({hidden})=0,"✓ All accounts listed","⚠ "&({hidden})&'
                                     f'" account(s) not listed individually — totals are still complete")')
    status_colors(ws, f"B{r}", f"B{r}")
    rows["CHECK"] = r

    names = {"Revenue": "Rev", "Cost of Sales": "COS", "GP": "GP", "Operating Expense": "Opex",
             "NOI": "NOI", "Other Income": "OthInc", "Other Expense": "OthExp", "NET": "Net", "GM": "GM", "NM": "NM"}
    for key, short in names.items():
        for col, suffix in [("C", "Cur"), ("D", "Prev"), ("E", "YTD"), ("F", "FY")]:
            define(wb, f"PL_{short}_{suffix}", ws.title, f"${col}${rows[key]}")
    ws.freeze_panes = "C7"
    _print_setup(ws, True, f"A1:S{r}")
    protect(ws)
    return rows


# ----------------------------------------------------------------------------- BALANCE SHEET
def build_balance_sheet(wb, ws) -> dict:
    _report_header(ws, "Balance Sheet",
                   '="As of "&TEXT(AsOf,"dd mmmm yyyy")&"  ·  Amounts in "&Currency', 3)
    ws.column_dimensions["B"].width = 44
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    hr = 6
    _colhead(ws, hr, 2, "")
    _colhead(ws, hr, 3, '="As of "&TEXT(AsOf,"dd-mmm-yyyy")')
    _colhead(ws, hr, 4, '="As of "&TEXT(PrevEnd,"dd-mmm-yyyy")')
    ws.row_dimensions[hr].height = 28
    rows = {}
    r = hr + 2

    def slots(typ, sign):
        nonlocal r
        start = r
        for k in range(1, M.BS_SLOTS[typ] + 1):
            c = ws.cell(row=r, column=2, value=f'=IFERROR(INDEX(calc_Name,MATCH("{typ}|{k}",calc_Key,0)),"")')
            c.alignment = S.INDENT2
            c.font = S.font()
            for col, rng in [("C", "calc_Bal"), ("D", "calc_BalPrev")]:
                cell = ws[f"{col}{r}"]
                cell.value = f'=IF($B{r}="","",{sign}*IFERROR(INDEX({rng},MATCH($B{r},calc_Name,0)),0))'
                cell.number_format = S.NUM
                cell.font = S.font()
            r += 1
        return start

    def subtotal(label, typ, sign, key):
        nonlocal r
        _total(ws, r, label, 3)
        ws.cell(row=r, column=2).alignment = S.INDENT1
        for col, rng in [("C", "calc_Bal"), ("D", "calc_BalPrev")]:
            ws[f"{col}{r}"].value = f'={sign}*SUMIFS({rng},calc_Type,"{typ}")'
        rows[key] = r
        r += 2

    def sub_label(text):
        nonlocal r
        c = ws.cell(row=r, column=2, value=text)
        c.font = S.font(bold=True, color=S.SLATE)
        c.alignment = S.INDENT1
        r += 1

    _section(ws, r, "ASSETS", 3); r += 1
    sub_label("Current Assets"); slots("Current Asset", 1); subtotal("Total Current Assets", "Current Asset", 1, "CA")
    sub_label("Fixed Assets"); slots("Fixed Asset", 1); subtotal("Total Fixed Assets", "Fixed Asset", 1, "FA")
    _total(ws, r, "TOTAL ASSETS", 3, strong=True)
    for col in "CD":
        ws[f"{col}{r}"].value = f'={col}{rows["CA"]}+{col}{rows["FA"]}'
    rows["TA"] = r
    r += 2
    _section(ws, r, "LIABILITIES", 3); r += 1
    sub_label("Current Liabilities"); slots("Current Liability", -1)
    subtotal("Total Current Liabilities", "Current Liability", -1, "CL")
    sub_label("Long-term Liabilities"); slots("Long-term Liability", -1)
    subtotal("Total Long-term Liabilities", "Long-term Liability", -1, "LTL")
    _total(ws, r, "TOTAL LIABILITIES", 3)
    for col in "CD":
        ws[f"{col}{r}"].value = f'={col}{rows["CL"]}+{col}{rows["LTL"]}'
    rows["TL"] = r
    r += 2
    _section(ws, r, "EQUITY", 3); r += 1
    slots("Equity", -1)
    c = ws.cell(row=r, column=2, value="Profit from earlier years")
    c.alignment, c.font = S.INDENT2, S.font()
    for col in "CD":
        c = ws[f"{col}{r}"]
        c.value = '=-SUMPRODUCT((calc_Stmt="PL")*(calc_Open+calc_PreFY))'
        c.number_format = S.NUM
        c.font = S.font()
    rows["PRIOR"] = r
    r += 1
    c = ws.cell(row=r, column=2, value="Current year profit (year to date)")
    c.alignment, c.font = S.INDENT2, S.font()
    ws[f"C{r}"].value = '=-SUMPRODUCT((calc_Stmt="PL")*calc_YTD)'
    ws[f"D{r}"].value = '=-SUMPRODUCT((calc_Stmt="PL")*(calc_YTD-calc_Cur))'
    for col in "CD":
        ws[f"{col}{r}"].number_format = S.NUM
        ws[f"{col}{r}"].font = S.font()
    rows["CY"] = r
    r += 1
    _total(ws, r, "TOTAL EQUITY", 3)
    for col in "CD":
        ws[f"{col}{r}"].value = f'={sign_eq(col)}+{col}{rows["PRIOR"]}+{col}{rows["CY"]}'
    rows["TE"] = r
    r += 2
    _total(ws, r, "TOTAL LIABILITIES + EQUITY", 3, strong=True)
    for col in "CD":
        ws[f"{col}{r}"].value = f'={col}{rows["TL"]}+{col}{rows["TE"]}'
    rows["TLE"] = r
    r += 2
    for col in "CD":
        ws[f"{col}{r}"].value = f'=ROUND({col}{rows["TA"]}-{col}{rows["TLE"]},2)'
        ws[f"{col}{r}"].number_format = S.NUM
        ws[f"{col}{r}"].font = S.font(size=9, color=S.MUTED)
    ws.cell(row=r, column=2, value="Difference (must be zero)").font = S.font(size=9, color=S.MUTED)
    rows["DIFF"] = r
    r += 1
    ws[f"B{r}"].value = (f'=IF(AND(C{rows["DIFF"]}=0,D{rows["DIFF"]}=0),"✓ Balanced: Assets = Liabilities + Equity",'
                         f'"⚠ Balance Sheet is out of balance by "&Currency&" "&TEXT(ABS(C{rows["DIFF"]}),"#,##0.00"))')
    ws[f"B{r}"].font = S.font(bold=True)
    status_colors(ws, f"B{r}", f"B{r}")
    rows["STATUS"] = r
    hidden = "+".join(f'MAX(0,COUNTIFS(calc_Type,"{t}",calc_Name,"<>")-{n})' for t, n in M.BS_SLOTS.items())
    r += 1
    ws.cell(row=r, column=2, value=f'=IF(({hidden})=0,"",'
                                   f'"⚠ "&({hidden})&" account(s) not listed individually — totals are still complete")')
    status_colors(ws, f"B{r}", f"B{r}")

    for key, short in [("TA", "TotalAssets"), ("TL", "TotalLiab"), ("TE", "TotalEquity"), ("TLE", "TotalLE"),
                       ("DIFF", "Diff"), ("PRIOR", "PriorProfit"), ("CY", "CYProfit"), ("CA", "CurAssets"),
                       ("FA", "FixedAssets"), ("CL", "CurLiab"), ("LTL", "LTLiab")]:
        define(wb, f"BS_{short}", ws.title, f"$C${rows[key]}")
        define(wb, f"BS_{short}_Prev", ws.title, f"$D${rows[key]}")
    define(wb, "BS_Status", ws.title, f"$B${rows['STATUS']}")
    ws.freeze_panes = "C7"
    _print_setup(ws, False, f"A1:D{r}")
    protect(ws)
    return rows


def sign_eq(col):
    rng = "calc_Bal" if col == "C" else "calc_BalPrev"
    return f'-SUMIFS({rng},calc_Type,"Equity")'


# ----------------------------------------------------------------------------- CASH FLOW
def build_cash_flow(wb, ws) -> dict:
    _report_header(ws, "Cash Flow", PERIOD_TEXT, 4)
    ws.column_dimensions["B"].width = 46
    for col in "CDE":
        ws.column_dimensions[col].width = 18
    hr = 6
    _colhead(ws, hr, 2, "")
    _colhead(ws, hr, 3, "This Month")
    _colhead(ws, hr, 4, "Year to Date")
    ws.row_dimensions[hr].height = 24
    per = {"C": 'j_Per,RptIdx', "D": 'j_Per,">=1",j_Per,"<="&RptIdx'}
    rows = {}
    r = hr + 2

    def put(label, formulas, key, indent=True, bold=False):
        nonlocal r
        c = ws.cell(row=r, column=2, value=label)
        c.alignment = S.INDENT1 if indent else S.LEFT
        c.font = S.font(bold=bold)
        for col in "CD":
            cell = ws[f"{col}{r}"]
            cell.value = "=" + formulas(col)
            cell.number_format = S.NUM
            cell.font = S.font(bold=bold)
        rows[key] = r
        r += 1

    def flows(group, positive):
        op = '">0"' if positive else '"<0"'
        def f(col):
            base = f'SUMIFS(j_CFMain,j_CFGroup,"{group}",j_CFMain,{op},{per[col]})'
            if group == "Operating":
                base += f'+SUMIFS(j_CFTax,j_CFTax,{op},{per[col]})'
            return base
        return f

    put("Cash at beginning of period",
        lambda col: 'SUMIFS(calc_BalPrev,calc_CF,"Cash")' if col == "C"
        else 'SUMIFS(calc_Open,calc_CF,"Cash")+SUMIFS(calc_PreFY,calc_CF,"Cash")', "BEGIN", indent=False, bold=True)
    r += 1
    for group, title, in_lbl, out_lbl, key in [
        ("Operating", "OPERATING ACTIVITIES — running the business",
         "Money received (sales, customer payments, other income)", "Money paid (suppliers, expenses, wages, taxes)", "OP"),
        ("Investing", "INVESTING ACTIVITIES — equipment and other long-term assets",
         "Money received from selling assets", "Money paid for equipment / assets", "INV"),
        ("Financing", "FINANCING ACTIVITIES — owners and lenders",
         "Money in from owners and loans", "Money out to owners and loan repayments", "FIN"),
    ]:
        _section(ws, r, title, 3); r += 1
        put(in_lbl, flows(group, True), key + "_IN")
        put(out_lbl, flows(group, False), key + "_OUT")
        _total(ws, r, f"Net cash from {group.lower()} activities", 3)
        for col in "CD":
            ws[f"{col}{r}"].value = f'={col}{rows[key + "_IN"]}+{col}{rows[key + "_OUT"]}'
        rows[key] = r
        r += 2
    _total(ws, r, "NET CHANGE IN CASH", 3, strong=True)
    for col in "CD":
        ws[f"{col}{r}"].value = f'={col}{rows["OP"]}+{col}{rows["INV"]}+{col}{rows["FIN"]}'
    rows["NET"] = r
    r += 1
    _total(ws, r, "CASH AT END OF PERIOD", 3, strong=True)
    for col in "CD":
        ws[f"{col}{r}"].value = f'={col}{rows["BEGIN"]}+{col}{rows["NET"]}'
    rows["END"] = r
    r += 2
    ws.cell(row=r, column=2, value="Cash per Balance Sheet (check)").font = S.font(size=9, color=S.MUTED)
    for col in "CD":
        c = ws[f"{col}{r}"]
        c.value = '=SUMIFS(calc_Bal,calc_CF,"Cash")'
        c.number_format = S.NUM
        c.font = S.font(size=9, color=S.MUTED)
    rows["ENDBS"] = r
    r += 1
    ws[f"B{r}"].value = (f'=IF(AND(ROUND(C{rows["END"]}-C{rows["ENDBS"]},2)=0,ROUND(D{rows["END"]}-D{rows["ENDBS"]},2)=0),'
                         f'"✓ Cash flow reconciles to the Balance Sheet","⚠ Cash flow does not reconcile — check entries")')
    ws[f"B{r}"].font = S.font(bold=True)
    status_colors(ws, f"B{r}", f"B{r}")
    rows["STATUS"] = r
    r += 2

    _section(ws, r, "CASH BY ACCOUNT", 4); r += 1
    for col, label in [("B", "Account"), ("C", "Start of month"), ("D", "End of month"), ("E", "Change")]:
        _colhead(ws, r, ws[f"{col}1"].column, label)
    r += 1
    for k in range(1, M.CASH_SLOTS + 1):
        ws[f"B{r}"].value = f'=IFERROR(INDEX(calc_Name,MATCH("Cash|{k}",calc_CashKey,0)),"")'
        ws[f"B{r}"].alignment = S.INDENT1
        ws[f"C{r}"].value = f'=IF($B{r}="","",INDEX(calc_BalPrev,MATCH($B{r},calc_Name,0)))'
        ws[f"D{r}"].value = f'=IF($B{r}="","",INDEX(calc_Bal,MATCH($B{r},calc_Name,0)))'
        ws[f"E{r}"].value = f'=IF($B{r}="","",D{r}-C{r})'
        for col in "CDE":
            ws[f"{col}{r}"].number_format = S.NUM
        r += 1

    for key, short in [("BEGIN", "Begin"), ("OP", "Op"), ("INV", "Inv"), ("FIN", "Fin"), ("NET", "Net"),
                       ("END", "End"), ("ENDBS", "EndBS"), ("OP_IN", "OpIn"), ("OP_OUT", "OpOut")]:
        define(wb, f"CF_{short}_Cur", ws.title, f"$C${rows[key]}")
        define(wb, f"CF_{short}_YTD", ws.title, f"$D${rows[key]}")
    define(wb, "CF_Status", ws.title, f"$B${rows['STATUS']}")
    _print_setup(ws, False, f"A1:E{r}")
    protect(ws)
    return rows


# ----------------------------------------------------------------------------- MONTHLY SUMMARY
MS_COLS = [
    ("Month", "Month", "mmm yyyy"),
    ("Rev", "Revenue", S.NUM), ("COS", "Cost of Sales", S.NUM), ("GP", "Gross Profit", S.NUM),
    ("Opex", "Operating Expenses", S.NUM), ("Other", "Other Income (net)", S.NUM), ("Net", "Net Profit", S.NUM),
    ("NM", "Net Margin", S.PCT), ("In", "Cash In", S.NUM), ("Out", "Cash Out", S.NUM),
    ("NetCash", "Net Cash Flow", S.NUM), ("Cash", "Ending Cash", S.NUM), ("AR", "Receivables", S.NUM),
    ("AP", "Payables", S.NUM),
]
MS_HEADER = 6


def build_monthly(wb, ws) -> dict:
    _report_header(ws, "Monthly Summary",
                   '="Fiscal year "&TEXT(FYStart,"mmm yyyy")&" – "&TEXT(FYEnd,"mmm yyyy")&"  ·  Amounts in "&Currency',
                   len(MS_COLS))
    L = {k: get_column_letter(i + 2) for i, (k, _h, _f) in enumerate(MS_COLS)}
    for i, (k, h, fmt) in enumerate(MS_COLS):
        _colhead(ws, MS_HEADER, i + 2, h)
        ws.column_dimensions[L[k]].width = 14
    ws.row_dimensions[MS_HEADER].height = 30
    first = MS_HEADER + 1
    for k in range(1, 13):
        r = MS_HEADER + k
        m, e = f"calc_M{k}", f"calc_E{k}"
        per = f"j_Per,{k}"
        f = {
            "Month": f"=EDATE(FYStart,{k - 1})",
            "Rev": f'=-SUMIFS({m},calc_Type,"Revenue")',
            "COS": f'=SUMIFS({m},calc_Type,"Cost of Sales")',
            "GP": f"={L['Rev']}{r}-{L['COS']}{r}",
            "Opex": f'=SUMIFS({m},calc_Type,"Operating Expense")',
            "Other": f'=-SUMIFS({m},calc_Type,"Other Income")-SUMIFS({m},calc_Type,"Other Expense")',
            "Net": f"={L['GP']}{r}-{L['Opex']}{r}+{L['Other']}{r}",
            "NM": f'=IF({L["Rev"]}{r}=0,"",{L["Net"]}{r}/{L["Rev"]}{r})',
            "In": f'=SUMIFS(j_CFMain,j_CFMain,">0",{per})+SUMIFS(j_CFTax,j_CFTax,">0",{per})',
            "Out": f'=SUMIFS(j_CFMain,j_CFMain,"<0",{per})+SUMIFS(j_CFTax,j_CFTax,"<0",{per})',
            "NetCash": f"={L['In']}{r}+{L['Out']}{r}",
            "Cash": f'=SUMIFS({e},calc_CF,"Cash")',
            "AR": f"=SUMIFS({e},calc_Name,acctAR)",
            "AP": f"=-SUMIFS({e},calc_Name,acctAP)",
        }
        for key, _h, fmt in MS_COLS:
            c = ws[f"{L[key]}{r}"]
            c.value = f[key]
            c.number_format = fmt
            c.font = S.font()
            c.border = S.Border(bottom=S.thin)
    last = MS_HEADER + 12
    tr = last + 1
    ws[f"B{tr}"].value = "Full year"
    for key, _h, fmt in MS_COLS[1:]:
        c = ws[f"{L[key]}{tr}"]
        if key == "NM":
            c.value = f'=IF({L["Rev"]}{tr}=0,"",{L["Net"]}{tr}/{L["Rev"]}{tr})'
        elif key in ("Cash", "AR", "AP"):
            c.value = f"={L[key]}{last}"
        else:
            c.value = f"=SUM({L[key]}{first}:{L[key]}{last})"
        c.number_format = fmt
    _total(ws, tr, "Full year", len(MS_COLS), strong=True)
    for key, _h, fmt in MS_COLS[1:]:
        ws[f"{L[key]}{tr}"].number_format = fmt
    # highlight the selected report month
    ws.conditional_formatting.add(f"B{first}:{L['AP']}{last}", FormulaRule(
        formula=[f"ROW()-{MS_HEADER}=RptIdx"], fill=S.fill("E0F2F1"), font=S.font(bold=True, color=S.NAVY)))
    ws[f"B{tr + 2}"].value = "The highlighted row is the report month selected on the Dashboard. Cash/AR/AP are month-end balances."
    ws[f"B{tr + 2}"].font = S.font(size=9, italic=True, color=S.MUTED)
    for key, _h, _f in MS_COLS:
        define(wb, f"ms_{key}", ws.title, f"${L[key]}${first}:${L[key]}${last}")
    ws.freeze_panes = f"C{first}"
    _print_setup(ws, True, f"A1:{L['AP']}{tr + 2}")
    protect(ws)
    return {"first": first, "last": last, "letters": L}


# ----------------------------------------------------------------------------- TRIAL BALANCE
def build_trial_balance(wb, ws) -> dict:
    _report_header(ws, "Trial Balance",
                   '="As of "&TEXT(AsOf,"dd mmmm yyyy")&"  ·  Amounts in "&Currency', 6)
    for col, w in {"B": 9, "C": 32, "D": 20, "E": 16, "F": 16, "G": 18}.items():
        ws.column_dimensions[col].width = w
    first = 12
    last = first + M.COA_ROWS - 1
    ws["C6"].value = "Total debits"
    ws["E6"].value = f"=SUM(E{first}:E{last})"
    ws["C7"].value = "Total credits"
    ws["F7"].value = f"=SUM(F{first}:F{last})"
    ws["C8"].value = "Journal integrity (debits − credits of every posting)"
    ws["G8"].value = "=ROUND(SUM(j_DrAmt)+SUM(j_TaxDr)-SUM(j_CrAmt)-SUM(j_TaxCr),2)"
    ws["C9"].value = "Status"
    ws["D9"].value = '=IF(AND(ROUND(E6-F7,2)=0,G8=0),"✓ BALANCED","⚠ CHECK DATA")'
    for ref in ("E6", "F7", "G8"):
        ws[ref].number_format = S.NUM
        ws[ref].font = S.font(bold=True)
    ws["D9"].font = S.font(bold=True, size=12)
    status_colors(ws, "D9", "D9")
    hr = first - 1
    for i, h in enumerate(["Code", "Account Name", "Type", "Debit", "Credit", "Net Balance (Dr+ / Cr−)"]):
        _colhead(ws, hr, 2 + i, h)
    from .engine import CL, C_FIRST
    for i in range(M.COA_ROWS):
        r = first + i
        cr = C_FIRST + i
        nm = f"CALC!${CL['Name']}${cr}"
        bal = f"CALC!${CL['Bal']}${cr}"
        ws[f"B{r}"].value = f'=IF({nm}="","",CALC!${CL["Code"]}${cr})'
        ws[f"C{r}"].value = f'={nm}'
        ws[f"D{r}"].value = f'=IF({nm}="","",CALC!${CL["Type"]}${cr})'
        ws[f"E{r}"].value = f'=IF({nm}="","",IF(ROUND({bal},2)>0,{bal},0))'
        ws[f"F{r}"].value = f'=IF({nm}="","",IF(ROUND({bal},2)<0,-{bal},0))'
        ws[f"G{r}"].value = f'=IF({nm}="","",{bal})'
        for col in "EFG":
            ws[f"{col}{r}"].number_format = S.NUM
        for col in "BCDEFG":
            ws[f"{col}{r}"].font = S.font()
            ws[f"{col}{r}"].border = S.Border(bottom=S.Side(style="hair", color=S.LINE))
    define(wb, "TB_Debit", ws.title, "$E$6")
    define(wb, "TB_Credit", ws.title, "$F$7")
    define(wb, "TB_Journal", ws.title, "$G$8")
    define(wb, "TB_Status", ws.title, "$D$9")
    ws.freeze_panes = f"B{first}"
    _print_setup(ws, False, f"A1:G{last}")
    protect(ws)
    return {"first": first}


# ----------------------------------------------------------------------------- LEDGER
LEDGER_FIRST = 16


def build_ledger(wb, ws, default_account="Bank Account") -> dict:
    _report_header(ws, "Account Ledger", '="Every posting to one account, in date order  ·  Amounts in "&Currency', 9)
    for col, w in {"B": 7, "C": 13, "D": 12, "E": 12, "F": 32, "G": 20, "H": 14, "I": 14, "J": 16, "K": 4,
                   "L": 12}.items():
        ws.column_dimensions[col].width = w
    ws["B6"].value = "Account"
    ws["B7"].value = "From"
    ws["B8"].value = "To"
    ws["D6"] = default_account
    S.as_input(ws["D6"])
    ws.merge_cells("D6:F6")
    S.as_input(ws["D7"], S.DATE)
    S.as_input(ws["D8"], S.DATE)
    add_list_validation(ws, "D6", "=lstAccounts", strict=False)
    add_date_validation(ws, "D7:D8")
    ws["G7"].value = "=IF(ISNUMBER(D7),D7,FYStart)"
    ws["G8"].value = "=IF(ISNUMBER(D8),D8,AsOf)"
    ws["H7"].value = "← blank = start of fiscal year"
    ws["H8"].value = "← blank = report date"
    for ref in ("G7", "G8"):
        S.as_auto(ws[ref], S.DATE)
    for ref in ("H7", "H8"):
        ws[ref].font = S.font(size=9, italic=True, color=S.MUTED)
    define(wb, "LedAcct", ws.title, "$D$6")
    define(wb, "LedFrom", ws.title, "$G$7")
    define(wb, "LedTo", ws.title, "$G$8")
    ws["I6"].value = '=IF(LedAcct="","",IFERROR("Normal balance: "&INDEX(calc_Normal,MATCH(LedAcct,calc_Name,0)),"⛔ Unknown account"))'
    ws["I6"].font = S.font(size=9, color=S.MUTED)
    status_colors(ws, "I6", "I6")
    nsign = 'IF(IFERROR(INDEX(calc_Normal,MATCH(LedAcct,calc_Name,0)),"Debit")="Credit",-1,1)'
    define(wb, "LedSign", ws.title, "$L$6")
    ws["L6"].value = "=" + nsign
    ws["L6"].font = S.font(size=8, color=S.MUTED)
    before = '"<"&LedFrom'
    opening = (f'=IF(LedAcct="",0,LedSign*(IFERROR(INDEX(calc_Open,MATCH(LedAcct,calc_Name,0)),0)'
               f'+SUMIFS(j_DrAmt,j_DrAcct,LedAcct,j_Date,{before})-SUMIFS(j_CrAmt,j_CrAcct,LedAcct,j_Date,{before})'
               f'+IF(LedAcct=acctTax,SUMIFS(j_TaxDr,j_Date,{before})-SUMIFS(j_TaxCr,j_Date,{before}),0)))')
    last = LEDGER_FIRST + M.LEDGER_LINES - 1
    summary = [
        ("Opening balance", opening),
        ("Total debits", f"=SUM(H{LEDGER_FIRST}:H{last})"),
        ("Total credits", f"=SUM(I{LEDGER_FIRST}:I{last})"),
        ("Closing balance", "=J10+LedSign*(J11-J12)"),
    ]
    for r, (lab, fml) in zip(range(10, 14), summary):
        ws[f"I{r}"].value = lab
        ws[f"J{r}"].value = fml
        ws[f"J{r}"].number_format = S.NUM
        ws[f"I{r}"].font = ws[f"J{r}"].font = S.font(bold=r in (10, 13))
    define(wb, "Led_Open", ws.title, "$J$10")
    define(wb, "Led_Close", ws.title, "$J$13")
    ws["B10"].value = (f'=COUNT(j_Key)&" entries"&IF(COUNT(j_Key)>{M.LEDGER_LINES},'
                       f'" — showing the first {M.LEDGER_LINES}; narrow the date range","")')
    ws["B10"].font = S.font(size=9, italic=True, color=S.MUTED)
    hr = LEDGER_FIRST - 1
    for i, h in enumerate(["#", "Date", "Source", "Source ID", "Description", "Customer / Vendor",
                           "Debit", "Credit", "Running Balance", "", "key"]):
        if h:
            _colhead(ws, hr, 2 + i, h)
    for k in range(1, M.LEDGER_LINES + 1):
        r = LEDGER_FIRST + k - 1
        key = f"$L{r}"
        n = f"MOD({key},{M.KEY_MULT})"
        ws[f"L{r}"].value = f'=IFERROR(SMALL(j_Key,{k}),"")'
        ws[f"B{r}"].value = f'=IF({key}="","",{k})'
        ws[f"C{r}"].value = f'=IF({key}="","",INDEX(j_Date,{n}))'
        ws[f"D{r}"].value = f'=IF({key}="","",INDEX(j_Source,{n}))'
        ws[f"E{r}"].value = f'=IF({key}="","",INDEX(j_ID,{n}))'
        ws[f"F{r}"].value = f'=IF({key}="","",INDEX(j_Desc,{n}))'
        ws[f"G{r}"].value = f'=IF({key}="","",INDEX(j_Party,{n}))'
        ws[f"H{r}"].value = (f'=IF({key}="","",IF(INDEX(j_DrAcct,{n})=LedAcct,INDEX(j_DrAmt,{n}),0)'
                             f'+IF(LedAcct=acctTax,INDEX(j_TaxDr,{n}),0))')
        ws[f"I{r}"].value = (f'=IF({key}="","",IF(INDEX(j_CrAcct,{n})=LedAcct,INDEX(j_CrAmt,{n}),0)'
                             f'+IF(LedAcct=acctTax,INDEX(j_TaxCr,{n}),0))')
        prev = "$J$10" if k == 1 else f"J{r - 1}"
        ws[f"J{r}"].value = f'=IF({key}="","",{prev}+LedSign*(H{r}-I{r}))' if k == 1 else \
            f'=IF({key}="","",N({prev})+LedSign*(H{r}-I{r}))'
        ws[f"C{r}"].number_format = S.DATE
        for col in "HIJ":
            ws[f"{col}{r}"].number_format = S.NUM
        for col in "BCDEFGHIJ":
            ws[f"{col}{r}"].font = S.font(size=9)
        ws[f"L{r}"].font = S.font(size=7, color="CBD5E1")
    ws.freeze_panes = f"B{LEDGER_FIRST}"
    _print_setup(ws, True, f"A1:J{last}")
    protect(ws)
    unlock(ws["D7"]); unlock(ws["D8"])
    return {}
