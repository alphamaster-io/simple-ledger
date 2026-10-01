"""Input sheets: SETUP, LISTS, COA, CUSTOMERS, VENDORS, TRANSACTIONS, AR, AP.

Every posting rule lives here, as same-row helper columns on the source
registers (Debit account / Debit amount / Credit account / ...). The JOURNAL
only stacks those columns; reports only aggregate the JOURNAL.
"""

from __future__ import annotations

from . import model as M
from . import styles as S
from .xl import (Col, Register, add_date_validation, add_list_validation, add_number_validation, define,
                 nav_link, protect, status_colors, title_block)

HEADER_ROW = 6  # registers: title 1-2, rule 3, legend 4, group labels 5, header 6, data 7+

DATE_MIN, DATE_MAX = "DATE(2000,1,1)", "DATE(2099,12,31)"


def _legend(ws, row=4):
    ws.cell(row=row, column=2, value="Tip: work top-to-bottom, one row per entry. "
            "Rows with ⛔ are NOT included in reports until fixed; ⚠ rows are included but worth a look.") \
        .font = S.font(size=9, italic=True, color=S.MUTED)


# ----------------------------------------------------------------------------- SETUP
SETUP_FIELDS = [
    # key, label, default, fmt, note, validation
    ("section", "BUSINESS DETAILS"),
    ("BizName", "Business name", "My Business", None, "Shown on every report header.", None),
    ("BizType", "Business type", "Retail Shop", None, "Pick the closest match.", "=lstBizTypes"),
    ("OwnerName", "Owner name", None, None, "", None),
    ("BizAddress", "Business address", None, None, "", None),
    ("BizEmail", "Email", None, None, "", None),
    ("BizPhone", "Phone", None, None, "", None),
    ("section", "ACCOUNTING SETTINGS"),
    ("StartDate", "Accounting start date", None, S.DATE,
     "The day you start recording in this workbook. Opening balances are as of this date.", "date"),
    ("FYStartMonthName", "Fiscal year starts in", "January", None,
     "Most small businesses use January (calendar year).", "=lstMonths"),
    ("Currency", "Currency (code or symbol)", "PHP", None, "Used in report headings, e.g. PHP, USD, ₱, $.", None),
    ("OpenCash", "Beginning cash on hand", 0, S.NUM,
     "Cash you physically had on the start date. Posts to 'Cash on Hand'.", "num"),
    ("OpenBank", "Beginning bank balance", 0, S.NUM,
     "Bank balance on the start date. Posts to 'Bank Account'. Other opening balances: COA sheet.", "num"),
    ("DefaultTerms", "Default payment terms (days)", M.DEFAULT_TERMS_DAYS, S.INT,
     "Used for invoice/bill due dates when you leave Due Date blank.", "num"),
    ("section", "TAX SETTINGS"),
    ("TaxOn", "Track VAT / sales tax separately?", "No", None,
     "Yes = the Tax column is posted to the Tax Payable account. No = tax columns are ignored.", "=lstYesNo"),
    ("TaxLabel", "Tax name", "VAT", None, "e.g. VAT, GST, Sales Tax. Used in labels only.", None),
]


def build_setup(wb, ws, setup: dict) -> None:
    title_block(ws, "Setup", "Fill in the yellow cells once. Everything else in the workbook reads from here.", 4)
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 70
    r = 5
    for f in SETUP_FIELDS:
        if f[0] == "section":
            r += 1 if r > 5 else 0
            c = ws.cell(row=r, column=2, value=f[1])
            c.font = S.font(bold=True, color=S.TEAL)
            r += 1
            continue
        key, label, default, fmt, note, val = f
        ws.cell(row=r, column=2, value=label).font = S.font(color=S.SLATE)
        cell = ws.cell(row=r, column=3, value=setup.get(key, default))
        S.as_input(cell, fmt)
        ws.cell(row=r, column=4, value=note).font = S.font(size=9, italic=True, color=S.MUTED)
        define(wb, key, ws.title, f"$C${r}")
        if val == "date":
            add_date_validation(ws, f"C{r}")
        elif val == "num":
            add_number_validation(ws, f"C{r}")
        elif val:
            add_list_validation(ws, f"C{r}", val)
        r += 1

    r += 1
    ws.cell(row=r, column=2, value="REPORT PERIOD  (automatic — choose the month on the Dashboard)").font = \
        S.font(bold=True, color=S.TEAL)
    nav_link(ws, f"D{r}", "DASHBOARD", "→ Change report month on the Dashboard", "C5")
    r += 1
    derived = [
        ("RptYear", "Report year", '=IF(ISNUMBER(RptYearIn),RptYearIn,YEAR(TODAY()))', S.INT),
        ("RptMonth", "Report month (1-12)",
         '=IFERROR(MATCH(RptMonthIn,lstMonths,0),MONTH(TODAY()))', S.INT),
        ("RptStart", "Report period start", "=DATE(RptYear,RptMonth,1)", S.DATE),
        ("AsOf", "Report period end (as-of date)", "=EOMONTH(RptStart,0)", S.DATE),
        ("FYStartMonth", "Fiscal year start month (1-12)", "=IFERROR(MATCH(FYStartMonthName,lstMonths,0),1)", S.INT),
        ("FYStart", "Fiscal year start", "=DATE(RptYear-IF(RptMonth<FYStartMonth,1,0),FYStartMonth,1)", S.DATE),
        ("FYEnd", "Fiscal year end", "=EDATE(FYStart,12)-1", S.DATE),
        ("FYYM", "Fiscal year start (year*12+month)", "=YEAR(FYStart)*12+MONTH(FYStart)", S.INT),
        ("RptIdx", "Report month number within fiscal year",
         "=YEAR(RptStart)*12+MONTH(RptStart)-FYYM+1", S.INT),
        ("PrevStart", "Previous month start", "=EDATE(RptStart,-1)", S.DATE),
        ("PrevEnd", "Previous month end", "=RptStart-1", S.DATE),
        ("TermsDays", "Default terms used", "=IF(ISNUMBER(DefaultTerms),DefaultTerms,30)", S.INT),
    ]
    for key, label, formula, fmt in derived:
        ws.cell(row=r, column=2, value=label).font = S.font(color=S.MUTED)
        S.as_auto(ws.cell(row=r, column=3, value=formula), fmt)
        define(wb, key, ws.title, f"$C${r}")
        r += 1
    protect(ws)


# ----------------------------------------------------------------------------- LISTS
LISTS_TOP = 6      # header row of the three upper tables
LISTS_LOWER = 20   # header row of the simple lists


def build_lists(wb, ws) -> None:
    title_block(ws, "Lists & Settings",
                "Dropdown sources and posting rules. Yellow cells may be edited; grey cells are automatic.", 18)
    hr = LISTS_TOP

    def header(row, col, labels, color=S.NAVY):
        for i, lab in enumerate(labels):
            c = ws.cell(row=row, column=col + i, value=lab)
            c.font = S.font(bold=True, color=S.WHITE)
            c.fill = S.fill(color)
            c.alignment = S.CENTER

    # Transaction types (B:G) ----------------------------------------------------
    ws.cell(row=hr - 1, column=2, value="TRANSACTION TYPES — how each type posts").font = S.font(bold=True, color=S.TEAL)
    header(hr, 2, ["Type", "Direction", "Other side", "Tax allowed", "How to use", "Other side account"])
    first = hr + 1
    for i, (t, d, src, tax, help_) in enumerate(M.TRANSACTION_TYPES):
        r = first + i
        for j, v in enumerate([t, d, src, tax, help_]):
            c = ws.cell(row=r, column=2 + j, value=v)
            c.font = S.font(size=9)
            c.alignment = S.WRAP if j == 4 else S.LEFT
        S.as_auto(ws.cell(row=r, column=7, value=f'=IF(D{r}="Category","",IFERROR(INDEX(sys_Acct,MATCH(D{r},sys_Role,0)),""))'))
    last = first + len(M.TRANSACTION_TYPES) - 1
    for name, col in [("tt_Type", "B"), ("tt_Dir", "C"), ("tt_Src", "D"), ("tt_Tax", "E"), ("tt_Other", "G")]:
        define(wb, name, ws.title, f"${col}${first}:${col}${last}")
    define(wb, "lstTypes", ws.title, f"$B${first}:$B${last}")

    # System accounts (I:L) -------------------------------------------------------
    ws.cell(row=hr - 1, column=9, value="SYSTEM ACCOUNTS — must match names on the COA sheet").font = \
        S.font(bold=True, color=S.TEAL)
    header(hr, 9, ["Role", "Used for", "Account name", "Check"])
    for i, (key, label, default) in enumerate(M.SYSTEM_ACCOUNTS):
        r = first + i
        ws.cell(row=r, column=9, value=key).font = S.font(size=9, color=S.MUTED)
        ws.cell(row=r, column=10, value=label).font = S.font(size=9)
        S.as_input(ws.cell(row=r, column=11, value=default))
        S.as_auto(ws.cell(row=r, column=12, value=f'=IF(ISNUMBER(MATCH(K{r},coa_Name,0)),"✓ OK","⛔ Not found on COA")'))
        define(wb, key, ws.title, f"$K${r}")
    sys_last = first + len(M.SYSTEM_ACCOUNTS) - 1
    define(wb, "sys_Role", ws.title, f"$I${first}:$I${sys_last}")
    define(wb, "sys_Acct", ws.title, f"$K${first}:$K${sys_last}")
    define(wb, "sys_Check", ws.title, f"$L${first}:$L${sys_last}")
    status_colors(ws, f"L{first}:L{sys_last}", f"L{first}")

    # Account types (N:R) -----------------------------------------------------------
    ws.cell(row=hr - 1, column=14, value="ACCOUNT TYPES").font = S.font(bold=True, color=S.TEAL)
    header(hr, 14, ["Account type", "Statement", "Sign", "Default normal", "Report section"])
    for i, row in enumerate(M.ACCOUNT_TYPES):
        r = first + i
        for j, v in enumerate(row):
            ws.cell(row=r, column=14 + j, value=v).font = S.font(size=9)
    at_last = first + len(M.ACCOUNT_TYPES) - 1
    for name, col in [("at_Type", "N"), ("at_Stmt", "O"), ("at_Sign", "P"), ("at_Normal", "Q"), ("at_Section", "R")]:
        define(wb, name, ws.title, f"${col}${first}:${col}${at_last}")

    # Simple lists ---------------------------------------------------------------------
    lr = LISTS_LOWER
    ws.cell(row=lr - 1, column=2, value="DROPDOWN LISTS").font = S.font(bold=True, color=S.TEAL)
    simple = [
        ("lstPayMethods", "Payment methods", M.PAYMENT_METHODS, 15, True),
        ("lstBizTypes", "Business types", M.BUSINESS_TYPES, 12, True),
        ("lstMonths", "Months", M.MONTHS, 12, False),
        ("lstYesNo", "Yes / No", M.YES_NO, 2, False),
        ("lstCF", "Cash flow sections", M.CF_SECTIONS, 4, False),
        ("lstNormal", "Normal balance", ["Debit", "Credit"], 2, False),
        ("lstBuckets", "Aging buckets", M.AGING_BUCKETS, 5, False),
    ]
    for j, (name, label, values, slots, editable) in enumerate(simple):
        col = 2 + j
        header(lr, col, [label], S.NAVY if editable else S.SLATE)
        for i in range(slots):
            c = ws.cell(row=lr + 1 + i, column=col, value=values[i] if i < len(values) else None)
            if editable:
                S.as_input(c)
            else:
                c.font = S.font(size=9)
        letter = c.column_letter
        define(wb, name, ws.title, f"${letter}${lr + 1}:${letter}${lr + slots}")

    # Automatic lists built from the COA ------------------------------------------------
    pay_col, act_col = 2 + len(simple) + 1, 2 + len(simple) + 2
    header(lr, pay_col, ["Payment accounts (auto)"], S.SLATE)
    header(lr, act_col, ["Active accounts (auto)"], S.SLATE)
    for k in range(1, 31):
        S.as_auto(ws.cell(row=lr + k, column=pay_col,
                          value=f'=IFERROR(INDEX(coa_Name,MATCH({k},coa_PayRank,0)),"")'))
    for k in range(1, M.COA_ROWS + 1):
        S.as_auto(ws.cell(row=lr + k, column=act_col,
                          value=f'=IFERROR(INDEX(coa_Name,MATCH({k},coa_ActRank,0)),"")'))
    pl, al = ws.cell(row=lr, column=pay_col).column_letter, ws.cell(row=lr, column=act_col).column_letter
    define(wb, "lstPayAccts", ws.title, f"${pl}${lr + 1}:${pl}${lr + 30}")
    define(wb, "lstAccounts", ws.title, f"${al}${lr + 1}:${al}${lr + M.COA_ROWS}")

    for col, w in {"B": 18, "C": 14, "D": 14, "E": 12, "F": 46, "G": 22, "H": 14, "I": 14, "J": 30, "K": 26,
                   "L": 18, "N": 20, "O": 10, "P": 6, "Q": 12, "R": 22}.items():
        ws.column_dimensions[col].width = w
    ws.column_dimensions[pl].width = 28
    ws.column_dimensions[al].width = 30
    protect(ws)


# ----------------------------------------------------------------------------- COA
def coa_register(ws) -> Register:
    cols = [
        Col("Code", "Code", 8, fmt=S.INT, name="coa_Code"),
        Col("Name", "Account Name", 30, name="coa_Name"),
        Col("Type", "Account Type", 20, name="coa_Type"),
        Col("Normal", "Normal Balance", 10, name="coa_Normal"),
        Col("CF", "Cash Flow Section", 12, name="coa_CF"),
        Col("Pay", "Payment Account?", 10, name="coa_Pay"),
        Col("Active", "Active?", 8, name="coa_Active"),
        Col("Opening", "Opening Balance", 15, fmt=S.NUM, name="coa_Open"),
        Col("Desc", "Description", 46),
        Col("Balance", "Balance at Report Date", 16, "auto", fmt=S.NUM, name="coa_Balance",
            formula='=IF({Name}="","",IFERROR(INDEX(calc_Bal,MATCH({Name},calc_Name,0))*IF({Normal}="Credit",-1,1),0))'),
        Col("Check", "Check", 34, "auto", name="coa_Check", formula=(
            '=IF(AND({Name}="",{Code}="",{Type}=""),"",'
            'IF({Name}="","⛔ Enter an account name",'
            'IF(COUNTIF(coa_Name,{Name})>1,"⛔ Duplicate account name",'
            'IF(ISNA(MATCH({Type},at_Type,0)),"⛔ Select an account type",'
            'IF(AND({Normal}<>"Debit",{Normal}<>"Credit"),"⛔ Normal balance must be Debit or Credit",'
            'IF(ISNA(MATCH({CF},lstCF,0)),"⚠ Select a cash flow section",'
            'IF(AND({Code}<>"",COUNTIF(coa_Code,{Code})>1),"⚠ Duplicate account code",'
            'IF(AND(N({Opening})<>0,INDEX(at_Stmt,MATCH({Type},at_Type,0))="PL"),'
            '"⚠ Opening balance on a profit & loss account",'
            '"✓ OK"))))))))')),
        Col("ActRank", "Active #", 7, "helper", name="coa_ActRank",
            formula='=IF(AND({Name}<>"",{Active}<>"No"),COUNTIFS({@Name}:{Name},"<>",{@Active}:{Active},"<>No"),"")'),
        Col("PayRank", "Payment #", 7, "helper", name="coa_PayRank",
            formula='=IF(AND({Name}<>"",{Active}<>"No",{Pay}="Yes"),'
                    'COUNTIFS({@Name}:{Name},"<>",{@Active}:{Active},"<>No",{@Pay}:{Pay},"Yes"),"")'),
    ]
    return Register(ws, cols, HEADER_ROW, M.COA_ROWS, "tblAccounts", "coa")


def build_coa(wb, ws, extra_accounts=None, opening=None) -> Register:
    title_block(ws, "Chart of Accounts",
                "Your categories. Rename, add or deactivate accounts here — reports pick them up automatically.", 12)
    _legend(ws)
    reg = coa_register(ws)
    opening = opening or {}
    rows = []
    for code, name, typ, normal, cf, pay, desc, *active in list(M.DEFAULT_ACCOUNTS) + [
            tuple(a) for a in (extra_accounts or [])]:
        rec = {"Code": code, "Name": name, "Type": typ, "Normal": normal, "CF": cf, "Pay": pay,
               "Active": active[0] if active else "Yes", "Desc": desc}
        if name == "Cash on Hand":
            rec["Opening"] = "=OpenCash"
        elif name == "Bank Account":
            rec["Opening"] = "=OpenBank"
        elif name == "Opening Balance Equity":
            rec["Opening"] = None
        elif name in opening:
            rec["Opening"] = opening[name]
        rows.append(rec)
    reg.write(wb, rows)
    # Linked / system opening balances are automatic: grey them out.
    oc = reg.col("Opening")
    for i, rec in enumerate(rows):
        if isinstance(rec.get("Opening"), str) or rec["Name"] == "Opening Balance Equity":
            S.as_auto(ws[f"{oc}{reg.first + i}"], S.NUM)
            if rec["Name"] == "Opening Balance Equity":
                ws[f"{oc}{reg.first + i}"].value = "(automatic)"
    add_list_validation(ws, reg.sqref("Type"), "=at_Type")
    add_list_validation(ws, reg.sqref("Normal"), "=lstNormal")
    add_list_validation(ws, reg.sqref("CF"), "=lstCF")
    add_list_validation(ws, reg.sqref("Pay"), "=lstYesNo")
    add_list_validation(ws, reg.sqref("Active"), "=lstYesNo")
    add_number_validation(ws, reg.sqref("Opening"), warning=True)
    status_colors(ws, reg.sqref("Check"), f"{reg.col('Check')}{reg.first}")
    ws["D4"].value = None
    return reg


# ----------------------------------------------------------------------------- CUSTOMERS / VENDORS
def party_register(ws, kind: str) -> Register:
    p = "cust" if kind == "customer" else "vend"
    src = "ar" if kind == "customer" else "ap"
    label = "Customer" if kind == "customer" else "Vendor"
    doc = "Invoiced" if kind == "customer" else "Billed"
    cols = [
        Col("Name", f"{label} Name", 28, name=f"{p}_Name"),
        Col("Contact", "Contact Person", 20),
        Col("Email", "Email", 26),
        Col("Phone", "Phone", 16),
        Col("Address", "Address", 30),
        Col("Terms", "Payment Terms (days)", 12, fmt=S.INT, name=f"{p}_Terms"),
        Col("Notes", "Notes", 26),
        Col("Total", f"Total {doc}", 15, "auto", fmt=S.NUM,
            formula=f'=IF({{Name}}="","",SUMIFS({src}_Amount,{src}_Party,{{Name}},{src}_Posted,1,{src}_Date,"<="&AsOf))'),
        Col("Paid", "Paid", 15, "auto", fmt=S.NUM,
            formula=f'=IF({{Name}}="","",SUMIFS({src}_Paid,{src}_Party,{{Name}}))'),
        Col("Outstanding", "Outstanding", 15, "auto", fmt=S.NUM, name=f"{p}_Outstanding",
            formula=f'=IF({{Name}}="","",SUMIFS({src}_Balance,{src}_Party,{{Name}}))'),
        Col("Overdue", "Overdue", 15, "auto", fmt=S.NUM,
            formula=f'=IF({{Name}}="","",SUMIFS({src}_Balance,{src}_Party,{{Name}},{src}_Status,"Overdue"))'),
        Col("Check", "Check", 22, "auto",
            formula=f'=IF({{Name}}="","",IF(COUNTIF({p}_Name,{{Name}})>1,"⚠ Duplicate name","✓ OK"))'),
        Col("TopKey", "Rank key", 8, "helper", name=f"{p}_TopKey",
            formula='=IF(AND({Name}<>"",N({Outstanding})>0),{Outstanding}+{n}/1000000,"")'),
    ]
    rows = M.CUSTOMER_ROWS if kind == "customer" else M.VENDOR_ROWS
    return Register(ws, cols, HEADER_ROW, rows, "tblCustomers" if kind == "customer" else "tblVendors", p)


def build_party(wb, ws, kind: str, records) -> Register:
    label = "Customers" if kind == "customer" else "Vendors (Suppliers)"
    title_block(ws, label, f"Add each {'customer' if kind == 'customer' else 'supplier'} once; "
                "then pick them from dropdowns. Balances are as of the report date.", 12)
    _legend(ws)
    reg = party_register(ws, kind)
    reg.write(wb, records)
    add_number_validation(ws, reg.sqref("Terms"), minimum=0)
    status_colors(ws, reg.sqref("Check"), f"{reg.col('Check')}{reg.first}")
    return reg


# ----------------------------------------------------------------------------- TRANSACTIONS
def tx_register(ws) -> Register:
    ids = [f"TX-{i:04d}" for i in range(1, M.TX_ROWS + 1)]
    err = (
        '=IF({Blank}=1,"",'
        'IF({Date}="","Missing transaction date",'
        'IF(NOT(ISNUMBER({Date})),"Invalid date",'
        f'IF(OR({{Date}}<{DATE_MIN},{{Date}}>{DATE_MAX}),"Invalid date",'
        'IF({Type}="","Select a type",'
        'IF({Dir}="","Unknown type: "&{Type},'
        'IF({Other}="","Select an account (Category)",'
        'IF(ISNA(MATCH({Other},coa_Name,0)),"Unknown account: "&{Other},'
        'IF({Money}="","Select Paid From / Deposited To",'
        'IF(ISNA(MATCH({Money},coa_Name,0)),"Unknown account: "&{Money},'
        'IF(NOT(ISNUMBER({Amount})),"Enter an amount",'
        'IF({Amount}=0,"Enter an amount",'
        'IF(ABS({TaxEff})>ABS({Amount}),"Tax is larger than the amount","")))))))))))))'
    )
    warn = (
        '=IF(OR({Blank}=1,{Err}<>""),"",'
        'IF(AND({ID}<>"",COUNTIF(tx_ID,{ID})>1),"Duplicate transaction ID",'
        'IF(AND({Amount}<0,{Type}<>"Adjustment"),"Negative amount — check this entry",'
        'IF(AND({Type}="Customer Payment",{Customer}=""),"Add the customer",'
        'IF(AND({Type}="Customer Payment",OR({Ref}="",COUNTIF(ar_No,{Ref}&"")=0)),'
        '"Invoice # in Reference not found on AR sheet",'
        'IF(AND({Type}="Bill Payment",{Vendor}=""),"Add the vendor",'
        'IF(AND({Type}="Bill Payment",OR({Ref}="",COUNTIFS(ap_No,{Ref}&"",ap_Party,{Vendor}&"")=0)),'
        '"Bill # + vendor not found on AP sheet",'
        'IF(AND(ISNUMBER({Tax}),N({Tax})<>0,{TaxEff}=0),"Tax ignored (tax tracking off, or not allowed for this type)",'
        'IF({Other}={Money}&"","Same account on both sides",'
        'IF(AND(ISNUMBER(StartDate),{Date}<StartDate),"Date is before the accounting start date",'
        'IF(AND({Type}<>"Adjustment",IFERROR(INDEX(coa_Pay,MATCH({Money}&"",coa_Name,0)),"")<>"Yes"),'
        '"Paid From / Deposited To is not a payment account","")))))))))))'
    )
    cols = [
        Col("ID", "Transaction ID", 11, name="tx_ID", static=ids),
        Col("Date", "Date", 12, fmt=S.DATE, name="tx_Date"),
        Col("Type", "Type", 16, name="tx_Type"),
        Col("Category", "Category (Account)", 26, name="tx_Cat"),
        Col("Desc", "Description", 30, name="tx_Desc"),
        Col("Customer", "Customer", 18, name="tx_Cust"),
        Col("Vendor", "Vendor", 18, name="tx_Vend"),
        Col("Method", "Payment Method", 13),
        Col("Ref", "Reference No. / Invoice # / Bill #", 14, name="tx_Ref"),
        Col("Money", "Paid From / Deposited To", 22, name="tx_Money"),
        Col("Amount", "Amount (incl. tax)", 14, fmt=S.NUM, name="tx_Amount"),
        Col("Tax", "Tax", 11, fmt=S.NUM, name="tx_Tax"),
        Col("Notes", "Notes", 20),
        Col("Net", "Net Amount", 13, "auto", fmt=S.NUM, name="tx_Net",
            formula='=IF({Posted}=1,{Amount}-{TaxEff},"")'),
        Col("AcctType", "Account Type", 16, "auto",
            formula='=IF({Other}="","",IFERROR(INDEX(coa_Type,MATCH({Other},coa_Name,0)),""))'),
        Col("DrAcct", "Debit Account", 20, "auto", name="tx_DrAcct",
            formula='=IF({Posted}=1,IF({Dir}="In",{Money}&"",{Other}),"")'),
        Col("CrAcct", "Credit Account", 20, "auto", name="tx_CrAcct",
            formula='=IF({Posted}=1,IF({Dir}="In",{Other},{Money}&""),"")'),
        Col("Check", "Check", 40, "auto", name="tx_Check",
            formula='=IF({Err}<>"","⛔ "&{Err}&" — not posted",IF({Warn}<>"","⚠ "&{Warn},IF({Posted}=1,"✓ OK","")))'),
        # helpers -----------------------------------------------------------------
        Col("Blank", "Blank?", 6, "helper",
            formula='=IF(AND({Date}="",{Type}="",{Category}="",{Amount}="",{Money}=""),1,0)'),
        Col("Dir", "Direction", 7, "helper", formula='=IFERROR(INDEX(tt_Dir,MATCH({Type},tt_Type,0)),"")'),
        Col("Other", "Other side", 18, "helper",
            formula='=IF({Type}="","",IFERROR(IF(INDEX(tt_Other,MATCH({Type},tt_Type,0))="",{Category}&"",'
                    'INDEX(tt_Other,MATCH({Type},tt_Type,0))),""))'),
        Col("TaxEff", "Tax used", 9, "helper", fmt=S.NUM,
            formula='=IF(OR(TaxOn<>"Yes",NOT(ISNUMBER({Tax}))),0,'
                    'IF(IFERROR(INDEX(tt_Tax,MATCH({Type},tt_Type,0)),"No")="Yes",{Tax},0))'),
        Col("Err", "Error", 18, "helper", name="tx_Err", formula=err),
        Col("Warn", "Warning", 18, "helper", name="tx_Warn", formula=warn),
        Col("Posted", "Posted", 6, "helper", name="tx_Posted", formula='=IF(AND({Blank}=0,{Err}=""),1,0)'),
        Col("DrAmt", "Debit Amt", 11, "helper", fmt=S.NUM, name="tx_DrAmt",
            formula='=IF({Posted}=1,IF({Dir}="In",{Amount},{Amount}-{TaxEff}),0)'),
        Col("CrAmt", "Credit Amt", 11, "helper", fmt=S.NUM, name="tx_CrAmt",
            formula='=IF({Posted}=1,IF({Dir}="In",{Amount}-{TaxEff},{Amount}),0)'),
        Col("TaxDr", "Tax Debit", 10, "helper", fmt=S.NUM, name="tx_TaxDr",
            formula='=IF(AND({Posted}=1,{Dir}="Out"),{TaxEff},0)'),
        Col("TaxCr", "Tax Credit", 10, "helper", fmt=S.NUM, name="tx_TaxCr",
            formula='=IF(AND({Posted}=1,{Dir}="In"),{TaxEff},0)'),
    ]
    return Register(ws, cols, HEADER_ROW, M.TX_ROWS, "tblTransactions", "tx")


def build_transactions(wb, ws, records) -> Register:
    title_block(ws, "Transactions",
                "Record every money movement here: sales, expenses, payments, transfers. One row = one transaction.", 18)
    _legend(ws)
    ws["H2"].value = "Example: 15-Jan-2026 | Money Out | Rent | Shop rent January | … | Bank Account | 15,000.00"
    ws["H2"].font = S.font(size=9, italic=True, color=S.MUTED)
    reg = tx_register(ws)
    reg.write(wb, records)
    add_date_validation(ws, reg.sqref("Date"))
    add_list_validation(ws, reg.sqref("Type"), "=lstTypes", prompt="Money In / Money Out cover most entries. "
                        "Use Customer Payment / Bill Payment for AR/AP.", title="Type")
    add_list_validation(ws, reg.sqref("Category"), "=lstAccounts", strict=False,
                        prompt="What was it for? (Revenue or expense account.) Leave blank for Customer/Bill "
                               "Payment and Owner Investment/Draw.", title="Category")
    add_list_validation(ws, reg.sqref("Customer"), "=cust_Name", strict=False)
    add_list_validation(ws, reg.sqref("Vendor"), "=vend_Name", strict=False)
    add_list_validation(ws, reg.sqref("Method"), "=lstPayMethods", strict=False)
    add_list_validation(ws, reg.sqref("Money"), "=lstPayAccts", strict=False,
                        prompt="Where the money came from / went to: Cash on Hand, Bank Account, Credit Card …",
                        title="Paid From / Deposited To")
    add_number_validation(ws, reg.sqref("Amount"))
    add_number_validation(ws, reg.sqref("Tax"), minimum=0)
    status_colors(ws, reg.sqref("Check"), f"{reg.col('Check')}{reg.first}")
    return reg


# ----------------------------------------------------------------------------- AR / AP
def doc_register(ws, kind: str) -> Register:
    ar = kind == "ar"
    p = "ar" if ar else "ap"
    party_key = "Party"
    party_list = "cust" if ar else "vend"
    pay_type = "Customer Payment" if ar else "Bill Payment"
    paid_amt = "tx_DrAmt" if ar else "tx_CrAmt"
    vendor_match = "" if ar else ',tx_Vend,{Party}'
    doc = "invoice" if ar else "bill"
    err = (
        '=IF({Blank}=1,"",'
        f'IF({{No}}="","Enter a{"n" if ar else ""} {doc} #",'
        + ('' if ar else 'IF({Party}="","Select a vendor",') +
        f'IF({{Date}}="","Missing {doc} date",'
        'IF(NOT(ISNUMBER({Date})),"Invalid date",'
        f'IF(OR({{Date}}<{DATE_MIN},{{Date}}>{DATE_MAX}),"Invalid date",'
        f'IF({{Account}}="","Select {"a revenue" if ar else "an expense"} account",'
        'IF(ISNA(MATCH({Account},coa_Name,0)),"Unknown account: "&{Account},'
        'IF(NOT(ISNUMBER({Amount})),"Enter an amount",'
        'IF({Amount}<=0,"Amount must be positive",'
        'IF({TaxEff}>{Amount},"Tax is larger than the amount",""))))))))))'
        + ('' if ar else ')')
    )
    dup = 'COUNTIF(ar_No,{No})>1' if ar else 'COUNTIFS(ap_No,{No},ap_Party,{Party})>1'
    acct_type = 'IFERROR(INDEX(coa_Type,MATCH({Account},coa_Name,0)),"")'
    type_warn = (f'IF(AND({acct_type}<>"Revenue",{acct_type}<>"Other Income"),"Account is not a revenue account",'
                 if ar else
                 f'IF(OR({acct_type}="Revenue",{acct_type}="Other Income",{acct_type}="Equity"),'
                 '"Account is not an expense or asset account",')
    warn = (
        '=IF(OR({Blank}=1,{Err}<>""),"",'
        f'IF({dup},"Duplicate {doc} #",'
        + ('IF({Party}="","Select a customer",' if ar else '') +
        'IF(N({Balance})<0,"Overpaid by "&TEXT(-{Balance},"#,##0.00"),'
        'IF(AND(ISNUMBER({Tax}),N({Tax})<>0,{TaxEff}=0),"Tax ignored (tax tracking is off)",'
        'IF(AND(ISNUMBER(StartDate),{Date}<StartDate),"Date is before the accounting start date",'
        + type_warn + '""))))))' + (')' if ar else '')
    )
    cols = [
        Col("No", "Invoice #" if ar else "Bill #", 12, name=f"{p}_No"),
        Col("Date", "Invoice Date" if ar else "Bill Date", 12, fmt=S.DATE, name=f"{p}_Date"),
        Col(party_key, "Customer" if ar else "Vendor", 22, name=f"{p}_Party"),
        Col("Desc", "Description", 28, name=f"{p}_Desc"),
        Col("Account", "Revenue Account" if ar else "Expense / Asset Account", 22, name=f"{p}_Account"),
        Col("Amount", "Amount (incl. tax)", 14, fmt=S.NUM, name=f"{p}_Amount"),
        Col("Tax", "Tax", 11, fmt=S.NUM),
        Col("Due", "Due Date (optional)", 12, fmt=S.DATE),
        Col("EffDue", "Due Date Used", 12, "auto", fmt=S.DATE, name=f"{p}_EffDue",
            formula=('=IF({Posted}=1,IF(ISNUMBER({Due}),{Due},{Date}+IF(ISNUMBER(IFERROR(INDEX('
                     f'{party_list}_Terms,MATCH({{Party}},{party_list}_Name,0)),"")),INDEX({party_list}_Terms,'
                     f'MATCH({{Party}},{party_list}_Name,0)),TermsDays)),"")')),
        Col("Paid", "Paid Amount", 13, "auto", fmt=S.NUM, name=f"{p}_Paid",
            formula=(f'=IF({{Posted}}=1,SUMIFS({paid_amt},tx_Type,"{pay_type}",tx_Ref,{{No}}&""'
                     f'{vendor_match},tx_Date,"<="&AsOf),"")')),
        Col("Balance", "Balance", 13, "auto", fmt=S.NUM, name=f"{p}_Balance",
            formula='=IF({Posted}=1,IF({Date}>AsOf,0,ROUND({Amount}-{Paid},2)),"")'),
        Col("Status", "Status", 10, "auto", name=f"{p}_Status",
            formula=('=IF({Posted}<>1,"",IF({Date}>AsOf,"Future",IF({Balance}<0,"Overpaid",IF({Balance}=0,"Paid",'
                     'IF(AsOf>{EffDue},"Overdue",IF({Paid}>0,"Partial","Unpaid"))))))')),
        Col("DaysOut", "Days Outstanding", 11, "auto", fmt=S.INT,
            formula='=IF(AND({Posted}=1,N({Balance})>0),AsOf-{Date},"")'),
        Col("DaysOver", "Days Overdue", 10, "auto", fmt=S.INT, name=f"{p}_DaysOver",
            formula='=IF({DaysOut}="","",MAX(0,AsOf-{EffDue}))'),
        Col("Bucket", "Aging Bucket", 10, "auto", name=f"{p}_Bucket",
            formula=('=IF({DaysOver}="","",IF({DaysOver}=0,"Current",IF({DaysOver}<=30,"1-30",'
                     'IF({DaysOver}<=60,"31-60",IF({DaysOver}<=90,"61-90","90+")))))')),
        Col("Check", "Check", 36, "auto", name=f"{p}_Check",
            formula='=IF({Err}<>"","⛔ "&{Err}&" — not posted",IF({Warn}<>"","⚠ "&{Warn},IF({Posted}=1,"✓ OK","")))'),
        # helpers
        Col("Blank", "Blank?", 6, "helper",
            formula='=IF(AND({No}="",{Date}="",{Party}="",{Amount}="",{Account}=""),1,0)'),
        Col("TaxEff", "Tax used", 9, "helper", fmt=S.NUM,
            formula='=IF(AND(TaxOn="Yes",ISNUMBER({Tax})),{Tax},0)'),
        Col("Err", "Error", 18, "helper", name=f"{p}_Err", formula=err),
        Col("Warn", "Warning", 18, "helper", name=f"{p}_Warn", formula=warn),
        Col("Posted", "Posted", 6, "helper", name=f"{p}_Posted", formula='=IF(AND({Blank}=0,{Err}=""),1,0)'),
        Col("DrAcct", "Debit Account", 16, "helper", name=f"{p}_DrAcct",
            formula='=IF({Posted}=1,acctAR,"")' if ar else '=IF({Posted}=1,{Account},"")'),
        Col("CrAcct", "Credit Account", 16, "helper", name=f"{p}_CrAcct",
            formula='=IF({Posted}=1,{Account},"")' if ar else '=IF({Posted}=1,acctAP,"")'),
        Col("DrAmt", "Debit Amt", 11, "helper", fmt=S.NUM, name=f"{p}_DrAmt",
            formula='=IF({Posted}=1,{Amount},0)' if ar else '=IF({Posted}=1,{Amount}-{TaxEff},0)'),
        Col("CrAmt", "Credit Amt", 11, "helper", fmt=S.NUM, name=f"{p}_CrAmt",
            formula='=IF({Posted}=1,{Amount}-{TaxEff},0)' if ar else '=IF({Posted}=1,{Amount},0)'),
        Col("TaxAmt", "Tax Amt", 10, "helper", fmt=S.NUM, name=f"{p}_TaxAmt",
            formula='=IF({Posted}=1,{TaxEff},0)'),
    ]
    rows = M.INV_ROWS if ar else M.BILL_ROWS
    return Register(ws, cols, HEADER_ROW, rows, "tblInvoices" if ar else "tblBills", p)


def build_doc(wb, ws, kind: str, records) -> Register:
    ar = kind == "ar"
    title_block(ws, "Accounts Receivable — Customer Invoices" if ar else "Accounts Payable — Supplier Bills",
                ("Sales you made on credit. When the customer pays, record a 'Customer Payment' on Transactions "
                 "with the invoice # in Reference." if ar else
                 "Bills you received but pay later. When you pay, record a 'Bill Payment' on Transactions "
                 "with the bill # in Reference and the same vendor."), 16)
    _legend(ws)
    reg = doc_register(ws, kind)
    reg.write(wb, records)
    add_date_validation(ws, reg.sqref("Date"))
    add_date_validation(ws, reg.sqref("Due"))
    add_list_validation(ws, reg.sqref("Party"), "=cust_Name" if ar else "=vend_Name", strict=False)
    add_list_validation(ws, reg.sqref("Account"), "=lstAccounts", strict=False)
    add_number_validation(ws, reg.sqref("Amount"))
    add_number_validation(ws, reg.sqref("Tax"), minimum=0)
    status_colors(ws, reg.sqref("Check"), f"{reg.col('Check')}{reg.first}")
    st = f"{reg.col('Status')}{reg.first}"
    from openpyxl.formatting.rule import FormulaRule
    ws.conditional_formatting.add(reg.sqref("Status"), FormulaRule(
        formula=[f'{st}="Overdue"'], font=S.font(bold=True, color=S.ERR_RED)))
    ws.conditional_formatting.add(reg.sqref("Status"), FormulaRule(
        formula=[f'{st}="Paid"'], font=S.font(color=S.OK_GREEN)))
    return reg
