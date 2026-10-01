"""Low-level helpers: named ranges, data validation, and the Register abstraction.

A *Register* is a data-entry sheet backed by an Excel Table: a title block,
a header row, N pre-sized rows, input columns the user types into, and
automatic formula columns. Column formulas are written as templates using
``{ColumnKey}`` tokens that resolve to same-row cell references, e.g.
``=IF({Date}="","",{Amount}-{TaxEff})`` -> ``=IF(B6="","",K6-W6)``.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Protection
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import styles as S

TOKEN = re.compile(r"\{(@?\w+)\}")


def define(wb, name: str, sheet: str, ref: str) -> None:
    """Workbook-scoped defined name, e.g. define(wb, "AsOf", "SETUP", "$C$30")."""
    wb.defined_names[name] = DefinedName(name, attr_text=f"{sheet}!{ref}")


def abs_range(col: str, r1: int, r2: int) -> str:
    return f"${col}${r1}:${col}${r2}"


def add_list_validation(ws, sqref: str, source: str, *, strict: bool = True, prompt: str | None = None,
                        title: str | None = None):
    dv = DataValidation(
        type="list", formula1=source, allow_blank=True,
        showErrorMessage=True, errorStyle="stop" if strict else "warning",
        errorTitle="Please pick from the list",
        error=("Choose a value from the dropdown." if strict else
               "This value is not in the list. Add it to its master list first, or press Yes to keep it."),
    )
    if prompt:
        dv.showInputMessage = True
        dv.promptTitle = (title or "Tip")[:32]
        dv.prompt = prompt[:255]
    ws.add_data_validation(dv)
    dv.add(sqref)
    return dv


def add_date_validation(ws, sqref: str):
    dv = DataValidation(
        type="date", operator="between", formula1="DATE(2000,1,1)", formula2="DATE(2099,12,31)",
        allow_blank=True, showErrorMessage=True, errorStyle="stop",
        errorTitle="Invalid date", error="Enter a real date, e.g. 15-Jan-2026.",
    )
    ws.add_data_validation(dv)
    dv.add(sqref)
    return dv


def add_number_validation(ws, sqref: str, *, minimum=None, warning=False):
    if minimum is None:
        dv = DataValidation(type="decimal", operator="between", formula1="-1E+12", formula2="1E+12")
    else:
        dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=str(minimum))
    dv.allow_blank = True
    dv.showErrorMessage = True
    dv.errorStyle = "warning" if warning else "stop"
    dv.errorTitle = "Numbers only"
    dv.error = "Enter a number without currency symbols."
    ws.add_data_validation(dv)
    dv.add(sqref)
    return dv


def status_colors(ws, rng: str, first_cell: str) -> None:
    """Colour status text: ⛔ red, ⚠ amber, ✓ green."""
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first_cell},1)="⛔"'],
                                                  font=S.font(bold=True, color=S.ERR_RED), fill=S.fill(S.ERR_FILL)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first_cell},1)="⚠"'],
                                                  font=S.font(color=S.WARN_AMBER), fill=S.fill(S.WARN_FILL)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first_cell},1)="✓"'],
                                                  font=S.font(color=S.OK_GREEN)))


def title_block(ws, title: str, subtitle: str = "", width_cols: int = 10) -> None:
    ws.sheet_view.showGridLines = False
    ws["A1"].value = None
    ws.column_dimensions["A"].width = 2
    c = ws["B1"]
    c.value = title
    c.font = S.font(size=16, bold=True, color=S.NAVY)
    ws.row_dimensions[1].height = 26
    if subtitle:
        s = ws["B2"]
        s.value = subtitle
        s.font = S.font(size=10, italic=True, color=S.MUTED)
    for col in range(2, 2 + width_cols):
        ws.cell(row=3, column=col).border = S.Border(top=S.Side(style="medium", color=S.TEAL))


def nav_link(ws, cell: str, target_sheet: str, label: str, target_cell: str = "A1") -> None:
    c = ws[cell]
    c.value = f'=HYPERLINK("#{target_sheet}!{target_cell}","{label}")'
    c.font = S.font(color=S.TEAL, bold=True, underline="single")


def protect(ws) -> None:
    """Protect without a password: stops accidental edits, never locks the owner out."""
    p = ws.protection
    p.sheet = True
    p.formatColumns = False
    p.formatRows = False
    p.selectLockedCells = False
    p.selectUnlockedCells = False


def unlock(cell) -> None:
    cell.protection = Protection(locked=False)


@dataclass
class Col:
    key: str
    header: str
    width: float = 12
    kind: str = "input"          # input | auto | helper
    formula: str | None = None   # template; {Key} -> same-row ref, {n} -> 1-based row index
    fmt: str | None = None
    name: str | None = None      # defined name for the column range
    static: list | None = None   # pre-filled values for input columns (e.g. IDs)


@dataclass
class Register:
    ws: object
    cols: list[Col]
    header_row: int
    rows: int
    table_name: str
    prefix: str
    letters: dict = field(default_factory=dict)

    def __post_init__(self):
        for i, c in enumerate(self.cols):
            self.letters[c.key] = get_column_letter(i + 2)  # column A is a gutter

    @property
    def first(self) -> int:
        return self.header_row + 1

    @property
    def last(self) -> int:
        return self.header_row + self.rows

    def col(self, key: str) -> str:
        return self.letters[key]

    def rng(self, key: str) -> str:
        c = self.letters[key]
        return abs_range(c, self.first, self.last)

    def sqref(self, key: str) -> str:
        c = self.letters[key]
        return f"{c}{self.first}:{c}{self.last}"

    def render(self, template: str, r: int) -> str:
        n = r - self.header_row

        def sub(m):
            k = m.group(1)
            if k == "n":
                return str(n)
            if k == "r":
                return str(r)
            if k.startswith("@"):  # absolute reference to the first data row (for running counts)
                col = self.letters[k[1:]]
                return f"${col}${self.first}"
            return f"{self.letters[k]}{r}"

        return TOKEN.sub(sub, template)

    def write(self, wb, data_rows: list[dict] | None = None, style_name="TableStyleLight1",
              group_helpers=True) -> None:
        ws = self.ws
        data_rows = data_rows or []
        hr = self.header_row
        # Group labels above the header: what to type vs what is automatic.
        first_auto = next((c for c in self.cols if c.kind != "input"), None)
        ws.cell(row=hr - 1, column=2, value="▼ YOUR ENTRIES  (white columns — type or pick from the dropdown)") \
            .font = S.font(bold=True, color=S.NAVY, size=9)
        if first_auto:
            ws.cell(row=hr - 1, column=self.cols.index(first_auto) + 2,
                    value="▼ AUTOMATIC  (grey columns — calculated, do not type)").font = S.font(
                bold=True, color=S.MUTED, size=9)
        for i, c in enumerate(self.cols):
            letter = self.letters[c.key]
            ws.column_dimensions[letter].width = c.width
            h = ws.cell(row=hr, column=i + 2, value=c.header)
            h.font = S.font(bold=True, color=S.WHITE)
            h.fill = S.fill(S.NAVY if c.kind == "input" else S.SLATE)
            h.alignment = S.CENTER
            if c.name:
                define(wb, c.name, ws.title, self.rng(c.key))
        ws.row_dimensions[hr].height = 30

        for idx in range(self.rows):
            r = self.first + idx
            rec = data_rows[idx] if idx < len(data_rows) else {}
            for i, c in enumerate(self.cols):
                cell = ws.cell(row=r, column=i + 2)
                if c.kind == "input":
                    if c.key in rec and rec[c.key] not in (None, ""):
                        cell.value = rec[c.key]
                    elif c.static is not None and idx < len(c.static):
                        cell.value = c.static[idx]
                    elif c.formula and c.key not in rec:
                        cell.value = self.render(c.formula, r)
                    cell.font = S.font(color=S.INPUT_FONT)
                else:
                    cell.value = self.render(c.formula, r) if c.formula else None
                    cell.font = S.font(color=S.SLATE if c.kind == "auto" else S.MUTED, size=9 if c.kind == "helper" else 10)
                    cell.fill = S.fill(S.AUTO_FILL)
                if c.fmt:
                    cell.number_format = c.fmt

        ref = f"B{hr}:{get_column_letter(len(self.cols) + 1)}{self.last}"
        t = Table(displayName=self.table_name, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name=style_name, showRowStripes=False)
        ws.add_table(t)
        ws.freeze_panes = ws.cell(row=self.first, column=3)

        if group_helpers:
            helpers = [self.letters[c.key] for c in self.cols if c.kind == "helper"]
            if helpers:
                ws.column_dimensions.group(helpers[0], helpers[-1], hidden=True, outline_level=1)
