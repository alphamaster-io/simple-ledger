"""Build the workbook.

    python -m simpleledger.build                      # blank product + sample-data edition into dist/
    python -m simpleledger.build --data FILE --out X  # any data set (used by tests)
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.workbook.properties import CalcProperties

from . import __version__, engine, inputs, reports, styles, ui

ROOT = Path(__file__).resolve().parents[2]
SAMPLE = ROOT / "src" / "sample-data" / "sunrise_bakery.json"

SHEET_ORDER = [
    "START_HERE", "SETUP", "DASHBOARD", "TRANSACTIONS", "AR", "AP", "CUSTOMERS", "VENDORS", "COA",
    "PNL", "BALANCE_SHEET", "CASH_FLOW", "MONTHLY_SUMMARY", "TRIAL_BALANCE", "LEDGER", "HELP",
    "LISTS", "JOURNAL", "CALC",
]
DATE_KEYS = {"Date", "Due", "StartDate"}


def _coerce(rec: dict) -> dict:
    out = {}
    for k, v in rec.items():
        if k in DATE_KEYS and isinstance(v, str) and v:
            try:
                v = dt.date.fromisoformat(v)
            except ValueError:
                pass  # leave text as typed — the workbook flags it as an invalid date
        out[k] = v
    return out


def build(out_path: str | Path, data: dict | None = None) -> Path:
    data = data or {}
    setup = _coerce(data.get("setup", {}))
    wb = Workbook()
    wb.remove(wb.active)
    # House font for every cell that has no explicit style.
    wb._fonts[0] = Font(name=styles.FONT_NAME, size=10, color=styles.SLATE)
    ws = {name: wb.create_sheet(name) for name in SHEET_ORDER}

    inputs.build_lists(wb, ws["LISTS"])
    inputs.build_setup(wb, ws["SETUP"], setup)
    inputs.build_coa(wb, ws["COA"], data.get("extra_accounts"), data.get("opening"))
    inputs.build_party(wb, ws["CUSTOMERS"], "customer", data.get("customers", []))
    inputs.build_party(wb, ws["VENDORS"], "vendor", data.get("vendors", []))
    inputs.build_transactions(wb, ws["TRANSACTIONS"], [_coerce(r) for r in data.get("transactions", [])])
    inputs.build_doc(wb, ws["AR"], "ar", [_coerce(r) for r in data.get("invoices", [])])
    inputs.build_doc(wb, ws["AP"], "ap", [_coerce(r) for r in data.get("bills", [])])
    engine.build_journal(wb, ws["JOURNAL"])
    engine.build_calc(wb, ws["CALC"])
    reports.build_pnl(wb, ws["PNL"])
    reports.build_balance_sheet(wb, ws["BALANCE_SHEET"])
    reports.build_cash_flow(wb, ws["CASH_FLOW"])
    ms = reports.build_monthly(wb, ws["MONTHLY_SUMMARY"])
    reports.build_trial_balance(wb, ws["TRIAL_BALANCE"])
    reports.build_ledger(wb, ws["LEDGER"], data.get("ledger_account", "Bank Account"))
    ui.build_start(wb, ws["START_HERE"])
    ui.build_dashboard(wb, ws["DASHBOARD"], ms)
    ui.build_help(wb, ws["HELP"])

    report = data.get("report", {})
    if report.get("month"):
        ws["DASHBOARD"]["C5"].value = report["month"]
    if report.get("year"):
        ws["DASHBOARD"]["F5"].value = report["year"]

    # Excel recalculates everything on open (openpyxl stores no cached values).
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.properties.title = "SimpleLedger — Small Business Accounting Workbook"
    wb.properties.creator = "SimpleLedger"
    wb.properties.version = __version__
    wb.active = 0
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def load(path: str | Path) -> dict:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--data", help="JSON data file (default: build blank + sample editions)")
    ap.add_argument("--out", help="output .xlsx path")
    args = ap.parse_args(argv)
    if args.data:
        print(build(args.out or "out.xlsx", load(args.data)))
        return
    dist = ROOT / "dist" / "standard"
    blank = {"setup": {}, "report": {}}
    print(build(dist / "SimpleLedger-Small-Business-Accounting.xlsx", blank))
    print(build(dist / "SimpleLedger-Sample-Business-Data.xlsx", load(SAMPLE)))


if __name__ == "__main__":
    main()
