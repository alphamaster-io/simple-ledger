"""User-facing sheets: START_HERE, DASHBOARD, HELP."""

from __future__ import annotations

from openpyxl.chart import BarChart, Reference

from . import __version__
from . import model as M
from . import styles as S
from .xl import add_list_validation, add_number_validation, define, protect, status_colors, title_block

PRODUCT = "SimpleLedger"
PROMISE = ("Know where your money goes, know how much you earn, and understand your business finances "
           "from one simple Excel workbook.")


def _button(ws, cell, sheet, label, target="A1", color=S.TEAL):
    c = ws[cell]
    c.value = f'=HYPERLINK("#{sheet}!{target}","{label}")'
    c.font = S.font(bold=True, color=S.WHITE)
    c.fill = S.fill(color)
    c.alignment = S.CENTER


# ----------------------------------------------------------------------------- START HERE
def build_start(wb, ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 18
    ws["B2"].value = PRODUCT
    ws["B2"].font = S.font(size=26, bold=True, color=S.NAVY)
    ws["B3"].value = "Small Business Accounting Workbook"
    ws["B3"].font = S.font(size=13, color=S.TEAL, bold=True)
    ws["B4"].value = PROMISE
    ws["B4"].font = S.font(size=11, italic=True, color=S.SLATE)
    ws["B5"].value = '="Workbook for: "&BizName'
    ws["B5"].font = S.font(size=10, color=S.MUTED)
    for col in "BCDEFGH":
        ws[f"{col}6"].border = S.Border(bottom=S.Side(style="medium", color=S.TEAL))

    ws["B8"].value = "QUICK START"
    ws["B8"].font = S.font(bold=True, color=S.TEAL)
    buttons = [("B9", "SETUP", "1 · Setup"), ("C9", "TRANSACTIONS", "2 · Record"), ("D9", "DASHBOARD", "3 · Dashboard"),
               ("E9", "PNL", "Profit & Loss"), ("F9", "BALANCE_SHEET", "Balance Sheet"), ("G9", "CASH_FLOW", "Cash Flow"),
               ("H9", "HELP", "Help")]
    for cell, sheet, label in buttons:
        _button(ws, cell, sheet, label, color=S.NAVY if cell in ("B9", "C9", "D9") else S.TEAL)
    ws.row_dimensions[9].height = 28
    more = [("B10", "AR", "Invoices (AR)"), ("C10", "AP", "Bills (AP)"), ("D10", "CUSTOMERS", "Customers"),
            ("E10", "VENDORS", "Vendors"), ("F10", "COA", "Accounts"), ("G10", "MONTHLY_SUMMARY", "Monthly"),
            ("H10", "LEDGER", "Ledger")]
    for cell, sheet, label in more:
        _button(ws, cell, sheet, label, color="5B7083")
    ws.row_dimensions[10].height = 22

    steps = [
        ("Read this page", "Two minutes. It tells you where everything is."),
        ("Complete SETUP", "Business name, start date, currency, beginning cash and bank balances."),
        ("Add customers and vendors", "Optional but recommended: CUSTOMERS and VENDORS sheets."),
        ("Record transactions", "Every sale, expense, payment and transfer — one row each on TRANSACTIONS."),
        ("Record credit sales and unpaid bills", "Invoices you will be paid for later → AR. Bills you will pay later → AP."),
        ("Review the DASHBOARD", "Pick the month at the top. Revenue, expenses, profit and cash update automatically."),
        ("Review monthly reports", "Profit & Loss, Balance Sheet, Cash Flow, Monthly Summary."),
        ("Fix warnings", "The Dashboard's Data Health box lists anything that needs attention (⛔ / ⚠)."),
        ("Export / share", "File → Print or Save as PDF on any report sheet. Reports are print-ready."),
    ]
    ws["B12"].value = "YOUR WORKFLOW"
    ws["B12"].font = S.font(bold=True, color=S.TEAL)
    for i, (head, text) in enumerate(steps):
        r = 13 + i
        ws[f"B{r}"].value = f"{i + 1}. {head}"
        ws[f"B{r}"].font = S.font(bold=True, color=S.NAVY)
        ws[f"D{r}"].value = text
        ws[f"D{r}"].font = S.font(color=S.SLATE)

    r = 23
    ws[f"B{r}"].value = "COLOUR GUIDE"
    ws[f"B{r}"].font = S.font(bold=True, color=S.TEAL)
    legend = [
        (S.INPUT_FILL, S.INPUT_FONT, "Yellow cell", "A setting you fill in (SETUP, LISTS, report filters)."),
        (S.WHITE, S.INPUT_FONT, "White table column", "Where you type your data (dark-blue header)."),
        (S.AUTO_FILL, S.SLATE, "Grey cell / column", "Automatic. Calculated for you — do not type here."),
        (S.ERR_FILL, S.ERR_RED, "⛔ Error", "The row is NOT included in reports until you fix it."),
        (S.WARN_FILL, S.WARN_AMBER, "⚠ Warning", "Included in reports, but please double-check."),
        (S.OK_FILL, S.OK_GREEN, "✓ OK", "All good."),
    ]
    for i, (bg, fg, label, text) in enumerate(legend):
        rr = r + 1 + i
        c = ws[f"B{rr}"]
        c.value = label
        c.fill = S.fill(bg)
        c.font = S.font(bold=True, color=fg)
        c.border = S.BOX
        ws[f"C{rr}"].value = text
        ws[f"C{rr}"].font = S.font()

    r = 31
    ws[f"B{r}"].value = "GOOD TO KNOW"
    ws[f"B{r}"].font = S.font(bold=True, color=S.TEAL)
    notes = [
        "Everything is calculated from three input sheets: TRANSACTIONS, AR and AP. Reports never need typing.",
        "No accounting knowledge needed: choose a Type (Money In / Money Out …) and a Category. Debits and credits are automatic.",
        "Don't delete or rename sheets. Don't type over grey cells. Most report sheets are protected (no password) to prevent accidents.",
        f"Capacity: {M.TX_ROWS:,} transactions, {M.INV_ROWS} invoices and {M.BILL_ROWS} bills per workbook — "
        "plenty for a typical small business year. Start a fresh copy each fiscal year.",
        "Keep a backup: save a dated copy at the end of every month (e.g. Books-2026-09.xlsx).",
        "Requires Microsoft Excel 2010 or later (Windows or Mac) or Microsoft 365. No macros.",
    ]
    for i, t in enumerate(notes):
        ws[f"B{r + 1 + i}"].value = "•  " + t
        ws[f"B{r + 1 + i}"].font = S.font()
    ws[f"B{r + 8}"].value = f"{PRODUCT} v{__version__}"
    ws[f"B{r + 8}"].font = S.font(size=8, color=S.MUTED)
    ws.sheet_properties.tabColor = S.TEAL
    protect(ws)


# ----------------------------------------------------------------------------- DASHBOARD
def _card(ws, top, col0, title, value_formula, sub_formula, fmt=S.NUM0, accent=S.NAVY, name=None):
    from openpyxl.utils import get_column_letter as gl
    c1, c3 = gl(col0), gl(col0 + 2)
    for r in range(top, top + 3):
        ws.merge_cells(f"{c1}{r}:{c3}{r}")
        for c in range(col0, col0 + 3):
            cell = ws.cell(row=r, column=c)
            cell.fill = S.fill(S.LIGHT)
            cell.border = S.Border(left=S.Side(style="thick", color=accent) if c == col0 else None)
    t = ws[f"{c1}{top}"]
    t.value = title
    t.font = S.font(size=9, bold=True, color=S.MUTED)
    t.alignment = S.Alignment(horizontal="left", vertical="bottom", indent=1)
    v = ws[f"{c1}{top + 1}"]
    v.value = value_formula
    v.font = S.font(size=18, bold=True, color=accent)
    v.number_format = fmt
    v.alignment = S.Alignment(horizontal="left", vertical="center", indent=1)
    s = ws[f"{c1}{top + 2}"]
    s.value = sub_formula
    s.font = S.font(size=9, color=S.SLATE)
    s.alignment = S.Alignment(horizontal="left", vertical="top", indent=1)
    ws.row_dimensions[top].height = 18
    ws.row_dimensions[top + 1].height = 30
    ws.row_dimensions[top + 2].height = 18
    if name:
        from .xl import define as _d
        _d(ws.parent, name, ws.title, f"${c1}${top + 1}")


def build_dashboard(wb, ws, ms) -> None:
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
    from openpyxl.utils import get_column_letter as gl
    for c in range(2, 14):
        ws.column_dimensions[gl(c)].width = 11.5
    ws.column_dimensions["N"].width = 3
    ws["B1"].value = "Financial Dashboard"
    ws["B1"].font = S.font(size=18, bold=True, color=S.NAVY)
    ws["B2"].value = "=BizName"
    ws["B2"].font = S.font(size=11, bold=True, color=S.SLATE)
    for c in range(2, 14):
        ws.cell(row=3, column=c).border = S.Border(top=S.Side(style="medium", color=S.TEAL))

    ws["B5"].value = "Report month"
    ws["B5"].font = S.font(bold=True)
    ws.merge_cells("C5:D5")
    S.as_input(ws["C5"])
    ws["E5"].value = "Year"
    ws["E5"].font = S.font(bold=True)
    ws["E5"].alignment = S.RIGHT
    S.as_input(ws["F5"], S.INT)
    add_list_validation(ws, "C5", "=lstMonths", prompt="Choose the month to report on. Blank = current month.",
                        title="Report month")
    add_number_validation(ws, "F5", minimum=2000)
    define(wb, "RptMonthIn", ws.title, "$C$5")
    define(wb, "RptYearIn", ws.title, "$F$5")
    ws["H5"].value = '="Period: "&TEXT(RptStart,"mmmm yyyy")&"   ·   Year to date from "&TEXT(FYStart,"dd mmm yyyy")&"   ·   "&Currency'
    ws["H5"].font = S.font(size=9, italic=True, color=S.MUTED)

    exp_cur = "PL_COS_Cur+PL_Opex_Cur+PL_OthExp_Cur"
    exp_ytd = "PL_COS_YTD+PL_Opex_YTD+PL_OthExp_YTD"
    ar_gl = "SUMIFS(calc_Bal,calc_Name,acctAR)"
    ap_gl = "-SUMIFS(calc_Bal,calc_Name,acctAP)"
    _card(ws, 7, 2, "REVENUE · THIS MONTH", "=PL_Rev_Cur", '="Year to date: "&TEXT(PL_Rev_YTD,"#,##0")',
          name="DB_Revenue")
    _card(ws, 7, 5, "EXPENSES · THIS MONTH", f"={exp_cur}", f'="Year to date: "&TEXT({exp_ytd},"#,##0")',
          accent=S.WARN_AMBER, name="DB_Expenses")
    _card(ws, 7, 8, "NET PROFIT · THIS MONTH", "=PL_Net_Cur", '="Year to date: "&TEXT(PL_Net_YTD,"#,##0;(#,##0)")',
          accent=S.OK_GREEN, name="DB_Net")
    _card(ws, 7, 11, "CASH BALANCE · MONTH END", "=CF_End_Cur",
          '="Change this month: "&TEXT(CF_Net_Cur,"+#,##0;-#,##0;0")', accent=S.TEAL, name="DB_Cash")
    _card(ws, 11, 2, "RECEIVABLES · CUSTOMERS OWE YOU", f"={ar_gl}",
          '="Overdue: "&TEXT(SUMIFS(ar_Balance,ar_Status,"Overdue"),"#,##0")', name="DB_AR")
    _card(ws, 11, 5, "PAYABLES · YOU OWE SUPPLIERS", f"={ap_gl}",
          '="Overdue: "&TEXT(SUMIFS(ap_Balance,ap_Status,"Overdue"),"#,##0")', accent=S.WARN_AMBER, name="DB_AP")
    _card(ws, 11, 8, "GROSS MARGIN · YEAR TO DATE", '=IF(PL_GM_YTD="","",PL_GM_YTD)',
          '="This month: "&IF(PL_GM_Cur="","-",TEXT(PL_GM_Cur,"0.0%"))', fmt=S.PCT, accent=S.OK_GREEN, name="DB_GM")
    _card(ws, 11, 11, "NET MARGIN · YEAR TO DATE", '=IF(PL_NM_YTD="","",PL_NM_YTD)',
          '="This month: "&IF(PL_NM_Cur="","-",TEXT(PL_NM_Cur,"0.0%"))', fmt=S.PCT, accent=S.TEAL, name="DB_NM")

    # Data health ----------------------------------------------------------------------
    ws["B15"].value = "DATA HEALTH"
    ws["B15"].font = S.font(bold=True, color=S.TEAL)
    errs = 'COUNTIF(tx_Err,"?*")+COUNTIF(ar_Err,"?*")+COUNTIF(ap_Err,"?*")'
    warns = 'COUNTIF(tx_Warn,"?*")+COUNTIF(ar_Warn,"?*")+COUNTIF(ap_Warn,"?*")'
    health = [
        ("Trial balance", "=TB_Status"),
        ("Balance sheet", '=IF(BS_Diff=0,"✓ Balanced","⚠ Out of balance by "&TEXT(BS_Diff,"#,##0.00"))'),
        ("Cash flow", '=IF(LEFT(CF_Status,1)="✓","✓ Reconciles","⚠ Does not reconcile")'),
        ("Entries with errors", f'=IF({errs}=0,"✓ None","⛔ "&{errs}&" not posted — see the Check column")'),
        ("Entries with warnings", f'=IF({warns}=0,"✓ None","⚠ "&{warns}&" to review")'),
        ("Receivables register", f'=IF(ROUND(SUM(ar_Balance)-{ar_gl},2)=0,"✓ Matches the ledger",'
                                 f'"⚠ Differs from ledger by "&TEXT(SUM(ar_Balance)-{ar_gl},"#,##0.00"))'),
        ("Payables register", f'=IF(ROUND(SUM(ap_Balance)-({ap_gl}),2)=0,"✓ Matches the ledger",'
                              f'"⚠ Differs from ledger by "&TEXT(SUM(ap_Balance)-({ap_gl}),"#,##0.00"))'),
        ("Setup", '=IF(COUNTIF(sys_Check,"⛔*")+COUNTIF(coa_Check,"⛔*")>0,"⛔ Fix the Chart of Accounts / Lists",'
                  f'IF(COUNT(tx_Date)>{int(M.TX_ROWS * 0.9)},"⚠ Transactions sheet is nearly full","✓ OK"))'),
    ]
    for i, (label, formula) in enumerate(health):
        r = 16 + i % 4
        lc, vc = ("B", "D") if i < 4 else ("H", "J")
        ws[f"{lc}{r}"].value = label
        ws[f"{lc}{r}"].font = S.font(color=S.SLATE)
        ws[f"{vc}{r}"].value = formula
        ws[f"{vc}{r}"].font = S.font(bold=True)
        status_colors(ws, f"{vc}{r}", f"{vc}{r}")
    define(wb, "DB_Errors", ws.title, "$D$19")
    define(wb, "DB_Warnings", ws.title, "$J$16")

    # Chart data block (O:S) ------------------------------------------------------------
    ws["P1"].value = "Chart data (automatic)"
    ws["P1"].font = S.font(bold=True, color=S.MUTED, size=9)
    for col, w in {"O": 2, "P": 26, "Q": 13, "R": 13, "S": 13}.items():
        ws.column_dimensions[col].width = w
    hdr = lambda r, labels: [setattr(ws.cell(row=r, column=16 + i, value=v), "font",
                                     S.font(bold=True, size=9, color=S.MUTED)) for i, v in enumerate(labels)]
    hdr(3, ["Month", "Revenue", "Expenses", "Net Profit"])
    L = ms["letters"]
    for k in range(12):
        r = 4 + k
        src = ms["first"] + k
        ws[f"P{r}"].value = f"=MONTHLY_SUMMARY!{L['Month']}{src}"
        ws[f"P{r}"].number_format = "mmm"
        ws[f"Q{r}"].value = f"=MONTHLY_SUMMARY!{L['Rev']}{src}"
        ws[f"R{r}"].value = f"=MONTHLY_SUMMARY!{L['Rev']}{src}-MONTHLY_SUMMARY!{L['Net']}{src}"
        ws[f"S{r}"].value = f"=MONTHLY_SUMMARY!{L['Net']}{src}"
        for col in "QRS":
            ws[f"{col}{r}"].number_format = S.NUM0
    hdr(17, ["Top expenses (YTD)", "Amount"])
    for k in range(1, M.TOP_N + 1):
        r = 17 + k
        ws[f"P{r}"].value = f'=IFERROR(INDEX(calc_Name,MATCH(LARGE(calc_TopKey,{k}),calc_TopKey,0)),"")'
        ws[f"Q{r}"].value = f'=IF(P{r}="",0,INDEX(calc_YTD,MATCH(P{r},calc_Name,0)))'
        ws[f"Q{r}"].number_format = S.NUM0
    hdr(24, ["Aging (days overdue)", "Receivables", "Payables"])
    for k, b in enumerate(M.AGING_BUCKETS):
        r = 25 + k
        ws[f"P{r}"].value = b
        ws[f"Q{r}"].value = f'=SUMIFS(ar_Balance,ar_Bucket,P{r})'
        ws[f"R{r}"].value = f'=SUMIFS(ap_Balance,ap_Bucket,P{r})'
        for col in "QR":
            ws[f"{col}{r}"].number_format = S.NUM0
    define(wb, "DB_ARAging", ws.title, "$Q$25:$Q$29")
    define(wb, "DB_APAging", ws.title, "$R$25:$R$29")

    # Charts ------------------------------------------------------------------------------
    cats = Reference(ws, min_col=16, min_row=4, max_row=15)
    c1 = BarChart()
    c1.type = "col"
    c1.title = "Revenue vs Expenses by Month"
    c1.add_data(Reference(ws, min_col=17, max_col=18, min_row=3, max_row=15), titles_from_data=True)
    c1.set_categories(cats)
    c1.height, c1.width = 7.5, 15.5
    c1.y_axis.numFmt = "#,##0"
    c1.y_axis.majorGridlines = None
    c1.legend.position = "b"
    _colors(c1, ["1F3A5F", "E9A23B"])
    ws.add_chart(c1, "B22")

    c2 = BarChart()
    c2.type = "col"
    c2.title = "Net Profit by Month"
    c2.add_data(Reference(ws, min_col=19, min_row=3, max_row=15), titles_from_data=True)
    c2.set_categories(cats)
    c2.height, c2.width = 7.5, 15.5
    c2.y_axis.numFmt = "#,##0"
    c2.legend = None
    c2.y_axis.majorGridlines = None
    _colors(c2, ["2A9D8F"])
    ws.add_chart(c2, "H22")

    c3 = BarChart()
    c3.type = "bar"
    c3.title = "Top Expenses (Year to Date)"
    c3.add_data(Reference(ws, min_col=17, min_row=18, max_row=17 + M.TOP_N), titles_from_data=False)
    c3.set_categories(Reference(ws, min_col=16, min_row=18, max_row=17 + M.TOP_N))
    c3.height, c3.width = 7.5, 15.5
    c3.legend = None
    c3.x_axis.scaling.orientation = "maxMin"
    c3.y_axis.numFmt = "#,##0"
    c3.y_axis.majorGridlines = None
    _colors(c3, ["E9A23B"])
    ws.add_chart(c3, "B38")

    c4 = BarChart()
    c4.type = "col"
    c4.title = "Receivables & Payables Aging"
    c4.add_data(Reference(ws, min_col=17, max_col=18, min_row=24, max_row=29), titles_from_data=True)
    c4.set_categories(Reference(ws, min_col=16, min_row=25, max_row=29))
    c4.height, c4.width = 7.5, 15.5
    c4.y_axis.numFmt = "#,##0"
    c4.y_axis.majorGridlines = None
    c4.legend.position = "b"
    _colors(c4, ["1F3A5F", "C2410C"])
    ws.add_chart(c4, "H38")

    # Tables under the charts --------------------------------------------------------------
    r0 = 54
    ws[f"B{r0}"].value = "CASH FLOW · THIS MONTH"
    ws[f"H{r0}"].value = "WHO OWES WHOM"
    for ref in (f"B{r0}", f"H{r0}"):
        ws[ref].font = S.font(bold=True, color=S.TEAL)
    cf = [("Cash at start of month", "=CF_Begin_Cur"), ("+ Operating activities", "=CF_Op_Cur"),
          ("+ Investing activities", "=CF_Inv_Cur"), ("+ Financing activities", "=CF_Fin_Cur"),
          ("Cash at end of month", "=CF_End_Cur")]
    for i, (lab, f) in enumerate(cf):
        r = r0 + 1 + i
        ws[f"B{r}"].value = lab
        ws[f"E{r}"].value = f
        ws[f"E{r}"].number_format = S.NUM
        bold = i in (0, 4)
        ws[f"B{r}"].font = ws[f"E{r}"].font = S.font(bold=bold)
        if i == 4:
            for col in "BCDE":
                ws[f"{col}{r}"].border = S.TOTAL_RULE
    due = M.DUE_SOON_DAYS
    owe = [
        ("Receivables due in the next 14 days", f'=SUMIFS(ar_Balance,ar_EffDue,">="&AsOf,ar_EffDue,"<="&(AsOf+{due}))'),
        ("Payables due in the next 14 days", f'=SUMIFS(ap_Balance,ap_EffDue,">="&AsOf,ap_EffDue,"<="&(AsOf+{due}))'),
    ]
    for i, (lab, f) in enumerate(owe):
        r = r0 + 1 + i
        ws[f"H{r}"].value = lab
        ws[f"L{r}"].value = f
        ws[f"L{r}"].number_format = S.NUM
    ws[f"H{r0 + 4}"].value = "Top customers owing you"
    ws[f"K{r0 + 4}"].value = "Top vendors you owe"
    for ref in (f"H{r0 + 4}", f"K{r0 + 4}"):
        ws[ref].font = S.font(bold=True, size=9, color=S.MUTED)
    for k in range(1, 4):
        r = r0 + 4 + k
        ws[f"H{r}"].value = f'=IFERROR(INDEX(cust_Name,MATCH(LARGE(cust_TopKey,{k}),cust_TopKey,0)),"")'
        ws[f"J{r}"].value = f'=IF(H{r}="","",INDEX(cust_Outstanding,MATCH(H{r},cust_Name,0)))'
        ws[f"K{r}"].value = f'=IFERROR(INDEX(vend_Name,MATCH(LARGE(vend_TopKey,{k}),vend_TopKey,0)),"")'
        ws[f"M{r}"].value = f'=IF(K{r}="","",INDEX(vend_Outstanding,MATCH(K{r},vend_Name,0)))'
        for col in "JM":
            ws[f"{col}{r}"].number_format = S.NUM0
        for col in "HJKM":
            ws[f"{col}{r}"].font = S.font(size=9)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = f"A1:M{r0 + 8}"
    ws.sheet_properties.tabColor = S.NAVY
    protect(ws)


def _colors(chart, colors, line=False):
    for s, color in zip(chart.series, colors):
        if line:
            s.graphicalProperties.line.solidFill = color
            s.graphicalProperties.line.width = 28000
            s.smooth = False
        else:
            s.graphicalProperties.solidFill = color
            s.graphicalProperties.line.solidFill = color


# ----------------------------------------------------------------------------- HELP
HELP = [
    ("How to start", [
        "1. Open SETUP and fill in the yellow cells: business name, accounting start date, currency, and how much "
        "cash and bank money you had on the start date.",
        "2. Optional: on the COA (Chart of Accounts) sheet, enter other opening balances — e.g. a loan you already "
        "owe, or equipment you already own. The 'Opening Balance Equity' account balances them automatically.",
        "3. Add your regular customers (CUSTOMERS) and suppliers (VENDORS).",
        "4. Start recording on TRANSACTIONS. Pick the month on the DASHBOARD to see results.",
    ]),
    ("How to record income (a sale paid immediately)", [
        "TRANSACTIONS → Type: Money In · Category: Product Sales (or Service Income) · Paid From / Deposited To: "
        "where the money went (Cash on Hand, Bank Account, E-Wallet) · Amount: total received.",
        "If you track VAT / sales tax (SETUP), type the tax portion in Tax. Example: 1,120 received incl. 120 VAT → "
        "Amount 1,120, Tax 120. Revenue is 1,000; 120 goes to Tax Payable.",
    ]),
    ("How to record a sale on credit (customer pays later)", [
        "AR sheet → new row: Invoice #, date, customer, revenue account, amount. Leave Due Date blank to use the "
        "customer's payment terms.",
        "When the customer pays: TRANSACTIONS → Type: Customer Payment · Customer · Reference: the invoice # · "
        "Deposited To: Bank/Cash · Amount paid. Leave Category blank. Partial payments are fine — record each one.",
    ]),
    ("How to record expenses", [
        "Paid now: TRANSACTIONS → Type: Money Out · Category: the expense (Rent, Utilities …) · Paid From: Cash, "
        "Bank or Credit Card · Amount.",
        "Paid later (supplier bill): AP sheet → bill #, date, vendor, expense account, amount. When you pay it: "
        "TRANSACTIONS → Type: Bill Payment · Vendor · Reference: bill # · Paid From · Amount.",
        "Buying equipment: Money Out with Category = Equipment. It goes to the Balance Sheet, not the P&L.",
        "Paying back a loan: Money Out with Category = Loans Payable (interest part: Interest Expense, on a separate row).",
    ]),
    ("How to record transfers and owner money", [
        "Moving money between your own accounts (e.g. depositing cash to the bank, paying the credit card from the "
        "bank): Type: Transfer · Paid From: where it left · Category: where it arrived.",
        "Owner puts in personal money: Type: Owner Investment · Deposited To: Bank/Cash. Category is automatic.",
        "Owner takes money for personal use: Type: Owner Draw · Paid From: Bank/Cash. This is not an expense.",
        "Loan received: Money In · Category: Loans Payable.",
    ]),
    ("How to add customers, suppliers and accounts", [
        "Type a new row at the bottom of CUSTOMERS, VENDORS or COA. Dropdowns update automatically.",
        "New account? Give it a unique name, a type (e.g. Operating Expense), Normal Balance and Cash Flow Section. "
        "It appears on the reports automatically.",
        "Stopped using an account? Set Active? = No to hide it from dropdowns. Don't delete accounts that have "
        "transactions.",
    ]),
    ("How to read the Profit & Loss", [
        "Revenue − Cost of Sales = Gross Profit (what you keep after the direct cost of what you sold).",
        "Gross Profit − Operating Expenses = Net Operating Income. ± Other income/expenses = NET PROFIT.",
        "Columns: this month, previous month, year to date, full year, and each month of the fiscal year.",
        "Profit is not the same as cash: credit sales count as revenue before the money arrives; equipment "
        "purchases use cash but are not expenses.",
    ]),
    ("How to read the Balance Sheet", [
        "What the business owns (Assets) = what it owes (Liabilities) + the owner's stake (Equity).",
        "Equity includes money you put in, minus drawings, plus all profits so far.",
        "The status line must say ✓ Balanced. If not, see 'How to fix warnings'.",
    ]),
    ("How to read the Cash Flow", [
        "Shows where cash came from and went, in three groups: Operating (day-to-day business), Investing "
        "(equipment and assets), Financing (owners and loans).",
        "Cash at the end = cash at the start + net change. It must match the cash on the Balance Sheet.",
        "Transfers between your own cash/bank accounts are not inflows or outflows — they cancel out.",
    ]),
    ("How to fix warnings", [
        "⛔ on a row: that row is NOT included in reports. Read the message in the Check column — e.g. 'Missing "
        "transaction date', 'Select an account', 'Enter an amount' — and fix the white cells on that row.",
        "⚠ on a row: included, but unusual — e.g. a negative amount, a Customer Payment whose invoice # can't be "
        "found, or a duplicate ID. Fix it if it's a mistake.",
        "Receivables/Payables register differs from the ledger: a Customer/Bill Payment probably has the wrong "
        "Reference, or someone used Category = Accounts Receivable/Payable on a normal Money In/Out row.",
        "Trial balance 'CHECK DATA' or Balance Sheet out of balance: a grey formula cell was overwritten. Copy the "
        "formula from the row above, or restore from your backup.",
    ]),
    ("What not to edit", [
        "Grey cells and grey columns; the JOURNAL and CALC sheets (the engine); the report sheets.",
        "Don't rename an account that already has transactions unless you also update those rows (the Check "
        "column will show 'Unknown account').",
        "Don't insert columns inside the input tables. Adding rows at the bottom is fine up to the capacity.",
    ]),
    ("Frequently asked questions", [
        "Do I need accounting knowledge? No. Pick a Type and a Category; the workbook does the debits and credits.",
        "Does it calculate my income tax? No. It produces the profit figure your accountant or tax form needs. "
        "VAT/sales tax collected and paid is tracked in Tax Payable if you switch it on.",
        "Can I use it for several businesses? Use one copy of the workbook per business.",
        "Does it work in Google Sheets? It is built and supported for Microsoft Excel. See the product page for "
        "the current compatibility list.",
        "Are there macros? No. It is a plain .xlsx file — safe to open, no security prompts.",
        "Can I print reports? Yes. Every report sheet is set up to print or Save as PDF on standard paper.",
        "A new fiscal year? Start a fresh copy, enter your closing balances from the old Balance Sheet as opening "
        "balances (cash and bank on SETUP, the rest on COA).",
    ]),
]


def build_help(wb, ws) -> None:
    title_block(ws, "Help", "Everything you need to use the workbook. No accounting background required.", 3)
    ws.column_dimensions["B"].width = 120
    r = 5
    for head, paras in HELP:
        ws[f"B{r}"].value = head
        ws[f"B{r}"].font = S.font(size=12, bold=True, color=S.NAVY)
        r += 1
        for p in paras:
            c = ws[f"B{r}"]
            c.value = p
            c.font = S.font()
            c.alignment = S.WRAP
            ws.row_dimensions[r].height = 15 * max(1, (len(p) // 110) + 1)
            r += 1
        r += 1
    ws.sheet_properties.tabColor = S.TEAL
    protect(ws)
