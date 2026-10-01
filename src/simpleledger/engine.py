"""Calculation engine: JOURNAL (all postings, stacked) and CALC (balances per account).

JOURNAL rows 1..TX_ROWS mirror TRANSACTIONS rows, the next INV_ROWS mirror AR,
the last BILL_ROWS mirror AP. INDEX(range, n) is used instead of direct cell
references so deleting a row in a source table never produces #REF!.
"""

from __future__ import annotations

from openpyxl.utils import get_column_letter

from . import model as M
from . import styles as S
from .xl import define, protect, title_block

J_HEADER = 6
J_FIRST = J_HEADER + 1
J_LAST = J_HEADER + M.JOURNAL_ROWS

# key, header, width, defined name, fmt
J_COLS = [
    ("n", "#", 6, "j_N", S.INT),
    ("Source", "Source", 12, "j_Source", None),
    ("ID", "Source ID", 12, "j_ID", None),
    ("Date", "Date", 12, "j_Date", S.DATE),
    ("Desc", "Description", 28, "j_Desc", None),
    ("Party", "Customer / Vendor", 18, "j_Party", None),
    ("Posted", "Posted", 7, "j_Posted", S.INT),
    ("DrAcct", "Debit Account", 22, "j_DrAcct", None),
    ("DrAmt", "Debit Amount", 13, "j_DrAmt", S.NUM),
    ("CrAcct", "Credit Account", 22, "j_CrAcct", None),
    ("CrAmt", "Credit Amount", 13, "j_CrAmt", S.NUM),
    ("TaxDr", "Tax Debit", 11, "j_TaxDr", S.NUM),
    ("TaxCr", "Tax Credit", 11, "j_TaxCr", S.NUM),
    ("Per", "FY Period", 8, "j_Per", S.INT),
    ("DrCash", "Dr is cash", 7, "j_DrCash", S.INT),
    ("CrCash", "Cr is cash", 7, "j_CrCash", S.INT),
    ("CFGroup", "Cash Flow Section", 12, "j_CFGroup", None),
    ("CFMain", "Cash Effect (main)", 13, "j_CFMain", S.NUM),
    ("CFTax", "Cash Effect (tax)", 12, "j_CFTax", S.NUM),
    ("Key", "Ledger Key", 14, "j_Key", "0"),
]
JL = {k: get_column_letter(i + 2) for i, (k, *_rest) in enumerate(J_COLS)}


def _src_formulas(src: str, n: int, r: int) -> dict:
    """Formulas pulling source row n into journal row r."""
    L = {k: f"{JL[k]}{r}" for k in JL}
    P = L["Posted"]
    if src == "tx":
        pairs = dict(ID="tx_ID", Date="tx_Date", Desc="tx_Desc", DrAcct="tx_DrAcct", DrAmt="tx_DrAmt",
                     CrAcct="tx_CrAcct", CrAmt="tx_CrAmt", TaxDr="tx_TaxDr", TaxCr="tx_TaxCr")
        party = f'=IF({P}=1,IF(INDEX(tx_Cust,{n})&""<>"",INDEX(tx_Cust,{n})&"",INDEX(tx_Vend,{n})&""),"")'
        label = "Transaction"
    else:
        pairs = dict(ID=f"{src}_No", Date=f"{src}_Date", Desc=f"{src}_Desc", DrAcct=f"{src}_DrAcct",
                     DrAmt=f"{src}_DrAmt", CrAcct=f"{src}_CrAcct", CrAmt=f"{src}_CrAmt")
        pairs["TaxDr" if src == "ap" else "TaxCr"] = f"{src}_TaxAmt"
        party = f'=IF({P}=1,INDEX({src}_Party,{n})&"","")'
        label = "Invoice" if src == "ar" else "Bill"
    f = {
        "Source": label,
        "Posted": f"=IFERROR(N(INDEX({src}_Posted,{n})),0)",
        "Party": party,
    }
    for k, rng in pairs.items():
        if k in ("DrAmt", "CrAmt", "TaxDr", "TaxCr"):
            f[k] = f"=IF({P}=1,INDEX({rng},{n}),0)"
        elif k == "Date":
            f[k] = f'=IF({P}=1,INDEX({rng},{n}),"")'
        else:
            f[k] = f'=IF({P}=1,INDEX({rng},{n})&"","")'
    for k in ("TaxDr", "TaxCr"):
        f.setdefault(k, 0)
    return f


def build_journal(wb, ws) -> None:
    title_block(ws, "Journal (engine — do not edit)",
                "Every posted entry from Transactions, AR and AP in one list. Reports read only this sheet.",
                len(J_COLS))
    ws["B4"].value = ("How to read: each row debits one account and credits another for the same total. "
                      "Tax (if tracked) posts to the Tax account. Rows with Posted = 0 are empty or have errors.")
    ws["B4"].font = S.font(size=9, italic=True, color=S.MUTED)
    for i, (k, h, w, name, fmt) in enumerate(J_COLS):
        col = get_column_letter(i + 2)
        ws.column_dimensions[col].width = w
        c = ws.cell(row=J_HEADER, column=i + 2, value=h)
        c.font = S.font(bold=True, color=S.WHITE)
        c.fill = S.fill(S.SLATE)
        c.alignment = S.CENTER
        define(wb, name, ws.title, f"${col}${J_FIRST}:${col}${J_LAST}")
    ws.row_dimensions[J_HEADER].height = 30

    regions = [("tx", M.TX_ROWS), ("ar", M.INV_ROWS), ("ap", M.BILL_ROWS)]
    r = J_FIRST
    jn = 1
    for src, count in regions:
        for n in range(1, count + 1):
            f = _src_formulas(src, n, r)
            L = {k: f"{JL[k]}{r}" for k in JL}
            P, D, H, J = L["Posted"], L["Date"], L["DrAcct"], L["CrAcct"]
            f["n"] = jn
            f["Per"] = f'=IF({P}=1,IF({D}<FYStart,0,IF({D}>FYEnd,13,YEAR({D})*12+MONTH({D})-FYYM+1)),"")'
            f["DrCash"] = f'=IF({P}=1,IF(IFERROR(INDEX(coa_CF,MATCH({H},coa_Name,0)),"")="Cash",1,0),0)'
            f["CrCash"] = f'=IF({P}=1,IF(IFERROR(INDEX(coa_CF,MATCH({J},coa_Name,0)),"")="Cash",1,0),0)'
            sign = f'({L["DrCash"]}-{L["CrCash"]})'
            f["CFGroup"] = (f'=IF({sign}=0,"",IFERROR(INDEX(coa_CF,MATCH(IF({sign}=1,{J},{H}),coa_Name,0)),'
                            f'"Operating"))')
            f["CFMain"] = f'=IF({sign}=1,{L["CrAmt"]},IF({sign}=-1,-{L["DrAmt"]},0))'
            f["CFTax"] = f'=IF({sign}=1,{L["TaxCr"]},IF({sign}=-1,-{L["TaxDr"]},0))'
            f["Key"] = (f'=IF(AND({P}=1,LedAcct<>""),IF(AND({D}>=LedFrom,{D}<=LedTo,OR({H}=LedAcct,{J}=LedAcct,'
                        f'AND(LedAcct=acctTax,{L["TaxDr"]}+{L["TaxCr"]}<>0))),{D}*{M.KEY_MULT}+{L["n"]},""),"")')
            for i, (k, _h, _w, _name, fmt) in enumerate(J_COLS):
                c = ws.cell(row=r, column=i + 2, value=f.get(k))
                c.font = S.font(size=9, color=S.SLATE)
                if fmt:
                    c.number_format = fmt
            r += 1
            jn += 1
    ws.freeze_panes = ws.cell(row=J_FIRST, column=3)
    ws.sheet_properties.tabColor = "94A3B8"
    protect(ws)


# ----------------------------------------------------------------------------- CALC
C_HEADER = 6
C_FIRST = C_HEADER + 1
C_LAST = C_HEADER + M.COA_ROWS
MONTH_KEYS = [f"M{k}" for k in range(1, 13)]
END_KEYS = [f"E{k}" for k in range(1, 13)]
C_COLS = (
    [("i", "#"), ("RawName", "COA name"), ("Name", "Account"), ("Code", "Code"), ("Type", "Type"),
     ("Normal", "Normal"), ("CF", "CF Section"), ("Stmt", "Statement"), ("Sign", "Sign"),
     ("RawOpen", "Opening (raw, Dr+)"), ("Open", "Opening (Dr+)"), ("PreFY", "Before FY (Dr+)")]
    + [(k, f"FY month {i + 1}") for i, k in enumerate(MONTH_KEYS)]
    + [("YTD", "Year to date"), ("Cur", "Report month"), ("Prev", "Previous month"), ("FY", "Full fiscal year"),
       ("Bal", "Balance at report date (Dr+)"), ("BalPrev", "Balance at prev. month end (Dr+)"),
       ("Key", "Report key"), ("CashKey", "Cash key"), ("TopKey", "Top expense key")]
    + [(k, f"Balance end of FY month {i + 1}") for i, k in enumerate(END_KEYS)]
)
CL = {k: get_column_letter(i + 2) for i, (k, _h) in enumerate(C_COLS)}


def _net(acct: str, *crit: str) -> str:
    """Net debit-minus-credit activity of an account over journal rows matching extra criteria."""
    c = "," + ",".join(crit) if crit else ""
    return (f"SUMIFS(j_DrAmt,j_DrAcct,{acct}{c})-SUMIFS(j_CrAmt,j_CrAcct,{acct}{c})"
            f"+IF({acct}=acctTax,SUMIFS(j_TaxDr,j_Posted,1{c})-SUMIFS(j_TaxCr,j_Posted,1{c}),0)")


def build_calc(wb, ws) -> None:
    title_block(ws, "Calc (engine — do not edit)",
                "Balances for every account and fiscal month. Positive = debit balance, negative = credit balance.",
                len(C_COLS))
    # Row 5: month index above the monthly columns (used by YTD SUMPRODUCT)
    for i, k in enumerate(MONTH_KEYS):
        c = ws[f"{CL[k]}5"]
        c.value = i + 1
        c.font = S.font(size=8, color=S.MUTED)
    for i, (k, h) in enumerate(C_COLS):
        col = CL[k]
        c = ws.cell(row=C_HEADER, column=i + 2, value=h)
        c.font = S.font(bold=True, color=S.WHITE, size=9)
        c.fill = S.fill(S.SLATE)
        c.alignment = S.CENTER
        ws.column_dimensions[col].width = 26 if k in ("Name", "RawName") else 12
        define(wb, f"calc_{k}", ws.title, f"${col}${C_FIRST}:${col}${C_LAST}")
    ws.row_dimensions[C_HEADER].height = 42
    m1, m12 = CL["M1"], CL["M12"]
    for i in range(1, M.COA_ROWS + 1):
        r = C_HEADER + i
        L = {k: f"{CL[k]}{r}" for k in CL}
        nm = f"${CL['Name']}{r}"
        f = {
            "i": i,
            "RawName": f'=IFERROR(INDEX(coa_Name,{i})&"","")',
            "Name": f'=IF({L["RawName"]}="","",IF(COUNTIF(${CL["RawName"]}${C_FIRST}:{L["RawName"]},{L["RawName"]})>1,"",'
                    f'{L["RawName"]}))',
            "Code": f'=IF({nm}="","",INDEX(coa_Code,{i}))',
            "Type": f'=IF({nm}="","",INDEX(coa_Type,{i})&"")',
            "Normal": f'=IF({nm}="","",INDEX(coa_Normal,{i})&"")',
            "CF": f'=IF({nm}="","",IF(INDEX(coa_CF,{i})&""="","Operating",INDEX(coa_CF,{i})&""))',
            "Stmt": f'=IF({nm}="","",IFERROR(INDEX(at_Stmt,MATCH({L["Type"]},at_Type,0)),""))',
            "Sign": f'=IF({nm}="",0,IFERROR(INDEX(at_Sign,MATCH({L["Type"]},at_Type,0)),0))',
            "RawOpen": f'=IF({nm}="",0,IF({L["Normal"]}="Credit",-1,1)*N(INDEX(coa_Open,{i})))',
            "Open": (f'=IF({nm}="",0,IF({nm}=acctOBE,{L["RawOpen"]}-SUM(${CL["RawOpen"]}${C_FIRST}:'
                     f'${CL["RawOpen"]}${C_LAST}),{L["RawOpen"]}))'),
            "PreFY": f'=IF({nm}="",0,{_net(nm, "j_Per,0")})',
            "YTD": f'=SUMPRODUCT((${m1}$5:${m12}$5<=RptIdx)*{m1}{r}:{m12}{r})',
            "Cur": f'=IF(AND(RptIdx>=1,RptIdx<=12),INDEX({m1}{r}:{m12}{r},RptIdx),0)',
            "FY": f"=SUM({m1}{r}:{m12}{r})",
            "Bal": f'={L["Open"]}+{L["PreFY"]}+{L["YTD"]}',
            "BalPrev": f'={L["Bal"]}-{L["Cur"]}',
            "Key": f'=IF({nm}="","",{L["Type"]}&"|"&COUNTIFS(${CL["Type"]}${C_FIRST}:{L["Type"]},{L["Type"]},'
                   f'${CL["Name"]}${C_FIRST}:{L["Name"]},"<>"))',
            "CashKey": (f'=IF(AND({nm}<>"",{L["CF"]}="Cash"),"Cash|"&COUNTIFS(${CL["CF"]}${C_FIRST}:{L["CF"]},"Cash",'
                        f'${CL["Name"]}${C_FIRST}:{L["Name"]},"<>"),"")'),
            "TopKey": (f'=IF(AND({nm}<>"",OR({L["Type"]}="Cost of Sales",{L["Type"]}="Operating Expense",'
                       f'{L["Type"]}="Other Expense"),{L["YTD"]}>0),{L["YTD"]}+{i}/1000000,"")'),
        }
        # Previous month: inside the FY use the monthly column; for FY month 1 use explicit dates.
        f["Prev"] = (f'=IF({nm}="",0,IF(RptIdx>1,INDEX({m1}{r}:{m12}{r},RptIdx-1),'
                     + _net(nm, 'j_Date,">="&PrevStart', 'j_Date,"<="&PrevEnd') + "))")
        for k_idx, k in enumerate(MONTH_KEYS):
            f[k] = f'=IF({nm}="",0,{_net(nm, f"j_Per,{k_idx + 1}")})'
        for k_idx, k in enumerate(END_KEYS):
            f[k] = f'={L["Open"]}+{L["PreFY"]}+SUM({m1}{r}:{CL[MONTH_KEYS[k_idx]]}{r})'
        for k, _h in C_COLS:
            c = ws[f"{CL[k]}{r}"]
            c.value = f[k]
            c.font = S.font(size=9, color=S.SLATE)
            if k not in ("i", "RawName", "Name", "Code", "Type", "Normal", "CF", "Stmt", "Sign", "Key", "CashKey"):
                c.number_format = S.NUM
    ws.freeze_panes = ws[f"{CL['Code']}{C_FIRST}"]
    ws.sheet_properties.tabColor = "94A3B8"
    protect(ws)
