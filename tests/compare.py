"""Compare a recalculated workbook with the reference model, figure by figure."""

from __future__ import annotations

import datetime as dt

from openpyxl.utils import get_column_letter

from reference import TYPE_INFO, Ledger, month_end

TOL = 0.011


class Mismatches(list):
    def check(self, label, got, want, tol=TOL):
        g = 0.0 if got in (None, "") else got
        w = 0.0 if want in (None, "") else want
        if isinstance(w, (int, float)) and not isinstance(w, bool):
            if not isinstance(g, (int, float)) or abs(g - w) > tol:
                self.append(f"{label}: workbook={got!r} reference={want!r}")
        elif g != w:
            self.append(f"{label}: workbook={got!r} reference={want!r}")


def _d(v):
    return v.date() if isinstance(v, dt.datetime) else v


def compare(book) -> tuple[Ledger, Mismatches]:
    ref = Ledger(book.data)
    mm = Mismatches()
    errs = book.errors()
    if errs:
        mm.append(f"formula errors: {errs[:10]}")

    # Journal: number of posted entries
    mm.check("posted entries", sum(v or 0 for v in book.name("j_Posted")), len(ref.postings), tol=0)

    # Account balances and openings
    calc = book.calc()
    for name in ref.accounts:
        row = calc.get(name)
        if row is None:
            mm.append(f"account missing from CALC: {name}")
            continue
        mm.check(f"opening {name}", row["Open"], ref.accounts[name]["open"])
        mm.check(f"balance {name}", row["Bal"], ref.balance(name))
        mm.check(f"balance prev {name}", row["BalPrev"], ref.balance(name, ref.prev_end))

    # P&L detail lines and totals
    ws = book.ws("PNL")
    labels = {ws[f"B{r}"].value: r for r in range(7, ws.max_row + 1) if ws[f"B{r}"].value}
    cols = [("C", "Cur"), ("D", "Prev"), ("E", "YTD"), ("F", "FY")] + [
        (get_column_letter(8 + k), f"M{k + 1}") for k in range(12)]
    for name, a in ref.accounts.items():
        if TYPE_INFO[a["type"]][0] != "PL":
            continue
        r = labels.get(name)
        if r is None:
            mm.append(f"P&L line missing: {name}")
            continue
        for col, per in cols:
            mm.check(f"P&L {name} {per}", ws[f"{col}{r}"].value, ref.pl_line(name, per))
    for short, typ in [("Rev", "Revenue"), ("COS", "Cost of Sales"), ("Opex", "Operating Expense"),
                       ("OthInc", "Other Income"), ("OthExp", "Other Expense")]:
        for per in ("Cur", "Prev", "YTD", "FY"):
            mm.check(f"PL_{short}_{per}", book.name(f"PL_{short}_{per}"), ref.pl_total(typ, per))
    for per in ("Cur", "Prev", "YTD", "FY"):
        mm.check(f"PL_Net_{per}", book.name(f"PL_Net_{per}"), ref.net_profit(per))

    # Balance sheet (this period and previous month end)
    for suffix, asof in (("", ref.asof), ("_Prev", ref.prev_end)):
        assets = sum(ref.bs_total(t, asof) for t in ("Current Asset", "Fixed Asset"))
        liab = sum(ref.bs_total(t, asof) for t in ("Current Liability", "Long-term Liability"))
        pl_all = -sum(ref.balance(a, asof) for a, v in ref.accounts.items() if TYPE_INFO[v["type"]][0] == "PL")
        equity = ref.bs_total("Equity", asof) + pl_all
        mm.check(f"BS_TotalAssets{suffix}", book.name(f"BS_TotalAssets{suffix}"), assets)
        mm.check(f"BS_TotalLiab{suffix}", book.name(f"BS_TotalLiab{suffix}"), liab)
        mm.check(f"BS_TotalEquity{suffix}", book.name(f"BS_TotalEquity{suffix}"), equity)
        mm.check(f"BS_Diff{suffix}", book.name(f"BS_Diff{suffix}"), 0)
    mm.check("BS_CYProfit", book.name("BS_CYProfit"), ref.net_profit("YTD"))
    if not str(book.name("BS_Status")).startswith("✓"):
        mm.append(f"BS status: {book.name('BS_Status')}")

    # Cash flow
    for suffix, lo, begin_date in (("Cur", ref.rpt_start, ref.prev_end),
                                    ("YTD", ref.fy_start, ref.fy_start - dt.timedelta(days=1))):
        flows = ref.cash_flows(lo, ref.asof)
        mm.check(f"CF_Begin_{suffix}", book.name(f"CF_Begin_{suffix}"), ref.cash_balance(begin_date))
        for short, group in (("Op", "Operating"), ("Inv", "Investing"), ("Fin", "Financing")):
            net = flows.get((group, "in"), 0) + flows.get((group, "out"), 0)
            mm.check(f"CF_{short}_{suffix}", book.name(f"CF_{short}_{suffix}"), net)
        mm.check(f"CF_OpIn_{suffix}", book.name(f"CF_OpIn_{suffix}"), flows.get(("Operating", "in"), 0))
        mm.check(f"CF_End_{suffix}", book.name(f"CF_End_{suffix}"), ref.cash_balance())
    if not str(book.name("CF_Status")).startswith("✓"):
        mm.append(f"CF status: {book.name('CF_Status')}")

    # Trial balance
    if book.name("TB_Status") != "✓ BALANCED":
        mm.append(f"TB status {book.name('TB_Status')}")
    mm.check("TB debits = credits", book.name("TB_Debit"), book.name("TB_Credit"))

    # AR / AP registers
    for sheet, docs in (("AR", ref.invoices), ("AP", ref.bills)):
        paid = book.column(sheet, "Paid Amount")
        bal = book.column(sheet, "Balance")
        status = book.column(sheet, "Status")
        bucket = book.column(sheet, "Aging Bucket")
        due = book.column(sheet, "Due Date Used")
        for i, doc in enumerate(docs):
            tag = f"{sheet} row {i + 1}"
            if doc is None:
                mm.check(f"{tag} status (not posted)", status[i], "")
                continue
            mm.check(f"{tag} paid", paid[i], doc.paid)
            mm.check(f"{tag} balance", bal[i], doc.balance)
            mm.check(f"{tag} status", status[i], doc.status)
            mm.check(f"{tag} bucket", bucket[i], doc.bucket)
            mm.check(f"{tag} due", _d(due[i]), doc.eff_due)

    # Dashboard
    mm.check("DB_AR", book.name("DB_AR"), ref.balance("Accounts Receivable"))
    mm.check("DB_AP", book.name("DB_AP"), -ref.balance("Accounts Payable"))
    mm.check("DB_Cash", book.name("DB_Cash"), ref.cash_balance())
    mm.check("DB_Revenue", book.name("DB_Revenue"), ref.pl_total("Revenue", "Cur"))
    mm.check("DB_Net", book.name("DB_Net"), ref.net_profit("Cur"))
    for name, docs in (("DB_ARAging", ref.invoices), ("DB_APAging", ref.bills)):
        want = [sum(d.balance for d in docs if d and d.bucket == b)
                for b in ("Current", "1-30", "31-60", "61-90", "90+")]
        for b, g, w in zip(("Current", "1-30", "31-60", "61-90", "90+"), book.name(name), want):
            mm.check(f"{name} {b}", g, w)
    n_err = len(ref.errors)
    db_err = book.name("DB_Errors")
    if n_err == 0:
        mm.check("DB_Errors", db_err, "✓ None")
    elif not str(db_err).startswith(f"⛔ {n_err} "):
        mm.append(f"DB_Errors: workbook={db_err!r} reference count={n_err}")

    # Monthly summary
    for k in range(1, 13):
        _lo, hi = ref.month_range(k)
        mm.check(f"ms revenue M{k}", book.name("ms_Rev")[k - 1], ref.pl_total("Revenue", f"M{k}"))
        mm.check(f"ms net M{k}", book.name("ms_Net")[k - 1], ref.net_profit(f"M{k}"))
        mm.check(f"ms cash M{k}", book.name("ms_Cash")[k - 1], ref.cash_balance(month_end(hi)))
        mm.check(f"ms AR M{k}", book.name("ms_AR")[k - 1], ref.balance("Accounts Receivable", hi))
        mm.check(f"ms AP M{k}", book.name("ms_AP")[k - 1], -ref.balance("Accounts Payable", hi))

    # Ledger (default: Bank Account, fiscal year start .. as-of)
    acct = book.name("LedAcct")
    if acct in ref.accounts:
        sign = -1 if ref.accounts[acct]["normal"] == "Credit" else 1
        mm.check("ledger opening", book.name("Led_Open"),
                 sign * ref.balance(acct, ref.fy_start - dt.timedelta(days=1)))
        mm.check("ledger closing", book.name("Led_Close"), sign * ref.balance(acct))
    return ref, mm
