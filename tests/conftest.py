"""Shared test fixtures.

Workbooks are built from data sets, recalculated once by LibreOffice (in
parallel), cached for the session, and exposed through the ``Book`` reader.
"""

from __future__ import annotations

import datetime as dt
import shutil
import sys
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "tools"))
sys.path.insert(0, str(Path(__file__).parent))

from simpleledger import build as B  # noqa: E402
from simpleledger import engine  # noqa: E402

import datasets  # noqa: E402
from recalc import recalc  # noqa: E402

HAS_LO = shutil.which("soffice") is not None
needs_lo = pytest.mark.skipif(not HAS_LO, reason="LibreOffice is required to recalculate workbooks")

ERROR_VALUES = {"#VALUE!", "#REF!", "#NAME?", "#DIV/0!", "#N/A", "#NUM!", "#NULL!"}


class Book:
    """Read-only view of a recalculated workbook."""

    def __init__(self, path: Path, data: dict):
        self.path = path
        self.data = data
        self.wb = load_workbook(path, data_only=True)

    def name(self, n: str):
        dn = self.wb.defined_names[n]
        (sheet, ref), = list(dn.destinations)
        ref = ref.replace("$", "")
        if ":" in ref:
            return [c.value for row in self.wb[sheet][ref] for c in row]
        return self.wb[sheet][ref].value

    def ws(self, sheet):
        return self.wb[sheet]

    def errors(self):
        found = []
        for ws in self.wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if isinstance(v, str) and (v in ERROR_VALUES or v.startswith("Err:")):
                        found.append(f"{ws.title}!{c.coordinate}={v}")
        return found

    def column(self, sheet, header, header_row=6):
        """Values of a register column (by header text) for all data rows."""
        ws = self.wb[sheet]
        for c in ws[header_row]:
            if c.value == header:
                col = c.column
                break
        else:
            raise KeyError(header)
        return [ws.cell(row=r, column=col).value for r in range(header_row + 1, ws.max_row + 1)]

    def row_by_label(self, sheet, label, col="B"):
        ws = self.wb[sheet]
        for r in range(1, ws.max_row + 1):
            if ws[f"{col}{r}"].value == label:
                return r
        raise KeyError(f"{label} not found on {sheet}")

    def calc(self):
        """{account name: {column key: value}} from the CALC sheet."""
        ws = self.wb["CALC"]
        out = {}
        for r in range(engine.C_FIRST, engine.C_LAST + 1):
            name = ws[f"{engine.CL['Name']}{r}"].value
            if name:
                out[name] = {k: ws[f"{engine.CL[k]}{r}"].value for k in engine.CL}
        return out


_CACHE: dict[str, Book] = {}


def _make(name: str, tmp: Path) -> Book:
    data = datasets.ALL[name]()
    raw = tmp / f"{name}.xlsx"
    B.build(raw, data)
    out = recalc(raw, tmp / f"{name}.calc.xlsx", timeout=600)
    return Book(out, data)


@pytest.fixture(scope="session")
def books(tmp_path_factory):
    """Dict-like access to every recalculated data set; all are built up front in parallel."""
    if not HAS_LO:
        pytest.skip("LibreOffice is required")
    tmp = tmp_path_factory.mktemp("books")
    missing = [n for n in datasets.ALL if n not in _CACHE]
    with ThreadPoolExecutor(max_workers=4) as pool:
        for name, book in zip(missing, pool.map(lambda n: _make(n, tmp), missing)):
            _CACHE[name] = book
    return _CACHE


def as_date(v):
    return v.date() if isinstance(v, dt.datetime) else v
