"""Static checks on the generated workbook (no LibreOffice needed)."""

import zipfile

import pytest
from openpyxl import load_workbook

from simpleledger import build as B
from simpleledger import model as M


@pytest.fixture(scope="module")
def blank_path(tmp_path_factory):
    return B.build(tmp_path_factory.mktemp("s") / "blank.xlsx", {"setup": {}, "report": {}})


@pytest.fixture(scope="module")
def wb(blank_path):
    return load_workbook(blank_path)


def test_sheets_in_order(wb):
    assert wb.sheetnames == B.SHEET_ORDER


def test_no_hidden_sheets(wb):
    assert all(ws.sheet_state == "visible" for ws in wb.worksheets)


def test_no_macros_or_external_links(blank_path):
    names = zipfile.ZipFile(blank_path).namelist()
    assert not any("vbaProject" in n or n.startswith("xl/externalLinks/") for n in names)


def test_tables_exist_with_expected_capacity(wb):
    tables = {t: ws.title for ws in wb.worksheets for t in ws.tables}
    assert tables == {"tblAccounts": "COA", "tblCustomers": "CUSTOMERS", "tblVendors": "VENDORS",
                      "tblTransactions": "TRANSACTIONS", "tblInvoices": "AR", "tblBills": "AP"}
    ref = wb["TRANSACTIONS"].tables["tblTransactions"].ref
    assert ref.endswith(str(6 + M.TX_ROWS))


def test_every_defined_name_points_to_a_real_sheet(wb):
    for name, dn in wb.defined_names.items():
        for sheet, _ref in dn.destinations:
            assert sheet in wb.sheetnames, name


def test_formulas_avoid_functions_with_poor_compatibility(wb):
    banned = ("XLOOKUP", "XMATCH", "FILTER(", "UNIQUE(", "SORT(", "SEQUENCE(", "LET(", "LAMBDA(",
              "INDIRECT(", "OFFSET(", "VSTACK(")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    up = c.value.upper()
                    assert not any(b in up for b in banned), f"{ws.title}!{c.coordinate}: {c.value[:80]}"


def test_report_sheets_are_protected_without_password(wb):
    for name in ("PNL", "BALANCE_SHEET", "CASH_FLOW", "TRIAL_BALANCE", "MONTHLY_SUMMARY", "JOURNAL", "CALC",
                 "DASHBOARD", "SETUP", "START_HERE", "HELP", "LEDGER", "LISTS"):
        p = wb[name].protection
        assert p.sheet and not p.password, name


def test_data_entry_sheets_are_not_protected(wb):
    for name in ("TRANSACTIONS", "AR", "AP", "CUSTOMERS", "VENDORS", "COA"):
        assert not wb[name].protection.sheet, name


def test_input_cells_on_protected_sheets_are_unlocked(wb):
    for ref in ("C5", "F5"):
        assert wb["DASHBOARD"][ref].protection.locked is False
    assert wb["LEDGER"]["D6"].protection.locked is False
    setup = wb["SETUP"]
    for name in ("BizName", "StartDate", "OpenCash", "TaxOn"):
        (sheet, ref), = list(wb.defined_names[name].destinations)
        assert setup[ref.replace("$", "")].protection.locked is False, name


def test_transaction_dropdowns_present(wb):
    ws = wb["TRANSACTIONS"]
    sources = {dv.formula1 for dv in ws.data_validations.dataValidation}
    assert {"lstTypes", "lstAccounts", "lstPayAccts", "cust_Name", "vend_Name"} <= {
        s.lstrip("=") for s in sources if s}


def test_default_chart_of_accounts_written(wb):
    names = [wb["COA"].cell(row=7 + i, column=3).value for i in range(len(M.DEFAULT_ACCOUNTS))]
    assert names == [a[1] for a in M.DEFAULT_ACCOUNTS]


def test_full_calc_on_load(wb):
    assert wb.calculation.fullCalcOnLoad


def test_dashboard_has_four_charts(wb):
    assert len(wb["DASHBOARD"]._charts) == 4
